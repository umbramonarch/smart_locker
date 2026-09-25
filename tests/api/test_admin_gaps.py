"""
File: test_admin_gaps.py
Description: Coverage-gap tests for admin and probe endpoints -- the Excel
             export download (xlsx body + headers), populated update-status
             marker merge, software-update launch and failure mapping,
             change-slot guards, registration gating (active session / no
             context), the public health payload, and the dry-run sync
             preview that must write nothing.
Project: smart_locker/tests/api
Notes: Run with: python -m pytest tests/api/test_admin_gaps.py -v
       BASE_DIR is patched per test so update-status.json / VERSION /
       update.sh fixtures live under tmp_path, never in the repo.
"""
import json
import subprocess
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

import smart_locker.api.app_context as ctx_module
import smart_locker.api.routes as routes_module
from smart_locker.database.models import Device, TransactionLog
from smart_locker.database.repositories import (
    DeviceRepository,
    RegistrantRepository,
)
from tests.api.helpers import catalog_workbook


def _enable_update_host(tmp_path, monkeypatch):
    """Point the update route at a fake host rooted at ``tmp_path``.

    routes.py reads ``BASE_DIR`` at call time for the update script path,
    the VERSION marker, and logs/update-status.json -- so one monkeypatch
    moves all three under the temp dir. ``_SYSTEMD_RUN`` just needs to be
    truthy to look like a Pi/systemd host.
    """
    script = tmp_path / "deploy" / "install" / "update.sh"
    script.parent.mkdir(parents=True)
    script.write_text("#!/bin/bash\nexit 0\n", encoding="utf-8")
    monkeypatch.setattr(routes_module, "BASE_DIR", tmp_path)
    monkeypatch.setattr(routes_module, "_SYSTEMD_RUN", "/usr/bin/systemd-run")
    return script


class TestExportExcelApi:
    """GET /api/admin/export-excel -- session + admin gate, real xlsx bytes."""

    def test_export_requires_session(self, client, mock_context):
        """No kiosk session -> 401 before any workbook is built."""
        resp = client.get("/api/admin/export-excel")
        assert resp.status_code == 401

    def test_export_rejects_non_admin(self, client, mock_context, test_user):
        """A normal user session cannot download the database export."""
        mock_context.session_mgr.start_session(test_user)
        resp = client.get("/api/admin/export-excel")
        assert resp.status_code == 403

    def test_export_admin_gets_valid_workbook(
        self, client, mock_context, admin_user, test_devices
    ):
        """Admin gets a real .xlsx attachment with the three export sheets."""
        mock_context.session_mgr.start_session(admin_user)
        resp = client.get("/api/admin/export-excel")
        assert resp.status_code == 200
        assert "spreadsheetml.sheet" in resp.headers["content-type"]
        disposition = resp.headers.get("content-disposition", "")
        assert "attachment" in disposition
        assert "smart_locker_data.xlsx" in disposition

        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames == ["Devices", "Transactions", "Users"]
        # Seeded devices land on the Devices sheet (header + one row each).
        devices_ws = wb["Devices"]
        assert devices_ws.max_row == 1 + len(test_devices)
        pms = {
            row[0] for row in devices_ws.iter_rows(min_row=2, values_only=True)
        }
        assert {d.pm_number for d in test_devices} <= pms
        # The admin user shows up on the Users sheet.
        users_ws = wb["Users"]
        names = {
            row[1] for row in users_ws.iter_rows(min_row=2, values_only=True)
        }
        assert "Admin User" in names


class TestUpdateStatusPopulated:
    """GET /api/admin/update-status merges the marker file and VERSION."""

    def test_update_status_returns_marker_fields(
        self, client, mock_context, admin_user, tmp_path, monkeypatch
    ):
        """state/message/at pass through; current_version comes from VERSION."""
        logs_dir = tmp_path / "logs"
        logs_dir.mkdir()
        (logs_dir / "update-status.json").write_text(
            json.dumps({
                "state": "success",
                "message": "Update applied cleanly.",
                "at": "2026-02-01T10:00:00+00:00",
                "version": "1.4.0",
            }),
            encoding="utf-8",
        )
        (tmp_path / "VERSION").write_text("1.4.0\n", encoding="utf-8")
        monkeypatch.setattr(routes_module, "BASE_DIR", tmp_path)

        mock_context.session_mgr.start_session(admin_user)
        resp = client.get("/api/admin/update-status")
        assert resp.status_code == 200
        body = resp.json()
        assert body["state"] == "success"
        assert body["message"] == "Update applied cleanly."
        assert body["at"] == "2026-02-01T10:00:00+00:00"
        assert body["version"] == "1.4.0"
        assert body["current_version"] == "1.4.0"


class TestTriggerUpdateLaunch:
    """POST /api/admin/update -- subprocess launch and failure mapping."""

    def test_update_launches_transient_unit(
        self, client, mock_context, admin_user, tmp_path, monkeypatch
    ):
        """Happy path: systemd-run is invoked and the endpoint reports started."""
        script = _enable_update_host(tmp_path, monkeypatch)
        calls = []

        def fake_run(cmd, **kwargs):
            calls.append((cmd, kwargs))
            return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

        monkeypatch.setattr(subprocess, "run", fake_run)
        mock_context.session_mgr.start_session(admin_user)
        resp = client.post("/api/admin/update")
        assert resp.status_code == 200
        body = resp.json()
        assert body["started"] is True
        assert "message" in body
        assert len(calls) == 1
        cmd, kwargs = calls[0]
        assert "systemd-run" in cmd
        assert "--unit=smart-locker-update" in cmd
        assert str(script) in cmd
        assert kwargs.get("check") is True
        assert kwargs.get("timeout") == 15

    def test_update_launch_failure_is_500(
        self, client, mock_context, admin_user, tmp_path, monkeypatch
    ):
        """systemd-run refusing (sudoers/cgroup) maps to 500 with stderr detail."""
        _enable_update_host(tmp_path, monkeypatch)

        def fake_run(cmd, **kwargs):
            raise subprocess.CalledProcessError(
                1, cmd, stderr="sudo: a password is required"
            )

        monkeypatch.setattr(subprocess, "run", fake_run)
        mock_context.session_mgr.start_session(admin_user)
        resp = client.post("/api/admin/update")
        assert resp.status_code == 500
        detail = resp.json()["detail"]
        assert "Could not start update" in detail
        assert "sudo: a password is required" in detail

    def test_update_launch_timeout_is_500(
        self, client, mock_context, admin_user, tmp_path, monkeypatch
    ):
        """A hung systemd-run hits the 15s timeout and still maps to 500."""
        _enable_update_host(tmp_path, monkeypatch)

        def fake_run(cmd, **kwargs):
            raise subprocess.TimeoutExpired(cmd, timeout=15)

        monkeypatch.setattr(subprocess, "run", fake_run)
        mock_context.session_mgr.start_session(admin_user)
        resp = client.post("/api/admin/update")
        assert resp.status_code == 500
        detail = resp.json()["detail"]
        assert "Could not start update" in detail
        assert "timed out" in detail


class TestSetDeviceSlotGaps:
    """POST /api/admin/devices/{id}/slot -- role/lookup/conflict guards."""

    def test_set_slot_rejects_non_admin(
        self, client, mock_context, test_user, test_devices
    ):
        """A logged-in non-admin cannot move a device between slots."""
        mock_context.session_mgr.start_session(test_user)
        resp = client.post(
            f"/api/admin/devices/{test_devices[0].id}/slot",
            json={"locker_slot": 9},
        )
        assert resp.status_code == 403

    def test_set_slot_unknown_device_404(
        self, client, mock_context, admin_user
    ):
        """Changing the slot on a missing device id is 404."""
        mock_context.session_mgr.start_session(admin_user)
        resp = client.post(
            "/api/admin/devices/99999/slot",
            json={"locker_slot": 9},
        )
        assert resp.status_code == 404

    def test_set_slot_conflict_is_409(
        self, client, mock_context, admin_user, test_devices, db_session
    ):
        """A slot already held by another device is 409 and changes nothing."""
        mock_context.session_mgr.start_session(admin_user)
        # Commit fixture rows first: the 409 path rolls back the shared
        # StaticPool connection, which would discard uncommitted fixtures.
        db_session.commit()
        occupied = test_devices[1].locker_slot  # Drone sits in slot 2
        resp = client.post(
            f"/api/admin/devices/{test_devices[0].id}/slot",
            json={"locker_slot": occupied},
        )
        assert resp.status_code == 409
        db_session.expire_all()
        assert test_devices[0].locker_slot == 1
        assert test_devices[1].locker_slot == occupied


class TestRegisterGatingGaps:
    """POST /api/register -- session conflict and missing-context guards."""

    def test_register_conflict_while_session_active(
        self, client, mock_context, test_user
    ):
        """A live kiosk session blocks self-registration with 409."""
        mock_context.session_mgr.start_session(test_user)
        resp = client.post("/api/register", json={"name": "Alice"})
        assert resp.status_code == 409
        assert mock_context.pending_registration is None

    def test_register_503_without_context(
        self, client, mock_context, monkeypatch
    ):
        """No AppContext at all -> 503 System not ready, not a 500."""
        monkeypatch.setattr(ctx_module, "context", None)
        resp = client.post("/api/register", json={"name": "Alice"})
        assert resp.status_code == 503


class TestHealthGaps:
    """GET /api/health -- public probe payload fields and session flag."""

    def test_health_payload_shape(self, client, mock_context):
        """Healthy DB probe -> ok; all probe keys present with sane types."""
        resp = client.get("/api/health")
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "ok"
        assert body["database"] is True
        assert isinstance(body["uptime_seconds"], (int, float))
        assert body["uptime_seconds"] >= 0
        assert body["session_active"] is False
        assert isinstance(body["nfc_reader"], bool)
        assert isinstance(body["fake_reader"], bool)
        assert "last_sync" in body
        assert "last_writeback" in body
        assert "update" in body

    def test_health_session_flag_flips(
        self, client, mock_context, test_user
    ):
        """session_active tracks the kiosk session manager state."""
        assert client.get("/api/health").json()["session_active"] is False
        mock_context.session_mgr.start_session(test_user)
        assert client.get("/api/health").json()["session_active"] is True
        mock_context.session_mgr.end_session()
        assert client.get("/api/health").json()["session_active"] is False


class TestSyncPreviewGaps:
    """POST /api/admin/sync-preview -- unset path and dry-run read-only."""

    def test_sync_preview_unconfigured_path_is_400(
        self, client, mock_context, admin_user
    ):
        """Empty SOURCE_EXCEL_PATH is a 400, not a crash or empty preview."""
        mock_context.session_mgr.start_session(admin_user)
        resp = client.post("/api/admin/sync-preview")
        assert resp.status_code == 400

    def test_sync_preview_dry_run_writes_nothing(
        self, client, mock_context, admin_user, test_devices, db_session,
        tmp_path, monkeypatch,
    ):
        """Dry-run reports would-change counts but leaves SQLite untouched."""
        path = catalog_workbook(tmp_path, [
            ["PM", "Name", "Location"],
            ["PM-001", "Camera Mk II", "Locker"],   # name differs -> updated
            ["PM-002", "Drone", "Locker"],          # identical -> unchanged
            ["PM-999", "Van kit", "Workshop"],      # Excel-only -> skipped
        ])
        monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))
        # Commit fixture rows first: the dry-run shares the StaticPool
        # connection and rolls its own UPDATE back -- a pending fixture
        # transaction would be discarded with it.
        db_session.commit()

        mock_context.session_mgr.start_session(admin_user)
        resp = client.post("/api/admin/sync-preview")
        assert resp.status_code == 200
        body = resp.json()
        assert body["preview"] is True
        assert body["imported"] == 0
        assert body["updated"] == 1
        assert body["unchanged"] == 1
        assert body["skipped"] == 1
        assert body["errors"] == 0

        # Nothing was written: device rows keep their original values.
        db_session.expire_all()
        camera = DeviceRepository.find_by_pm(db_session, "PM-001")
        assert camera.name == "Camera"
        devices = db_session.execute(select(Device)).scalars().all()
        assert len(devices) == len(test_devices)
        assert db_session.execute(
            select(TransactionLog)
        ).scalars().all() == []
        # The Location column's person names are not synced in dry-run.
        assert RegistrantRepository.get_all(db_session) == []
