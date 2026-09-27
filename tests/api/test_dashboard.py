"""
File: test_dashboard.py
Description: Tests for public Inventory/Locker/Display GETs, the public
             owner edit on non-cabinet rows, catalog CRUD, mirror endpoints,
             and dashboard NFC bind/unbind.
Project: smart_locker/tests/api
Notes: Run with: python -m pytest tests/api/test_dashboard.py -v
"""
import asyncio
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient

from smart_locker.api.app_context import PendingRegistration, PendingTagBind
from smart_locker.api.routes import router
from smart_locker.auth.session_manager import SessionManager
from smart_locker.database.models import DeviceStatus, UserRole
from smart_locker.database.repositories import DeviceRepository, RegistrantRepository, UserRepository
from smart_locker.security.hashing import compute_uid_hmac

import smart_locker.api.app_context as ctx_module
from tests.api.helpers import dashboard_admin_headers


@pytest.fixture()
def catalog_device(db_session):
    """One non-locker catalog row (locker_slot NULL, stored place)."""
    device = DeviceRepository.create(
        db_session,
        name="Van kit",
        device_type="Tool",
        pm_number="PM-999",
        manufacturer="Fluke",
        model="87V",
        serial_number="SN-9",
    )
    device.location = "Workshop"
    db_session.commit()
    return device


class TestDashboardInventoryAndDisplay:
    """Inventory and Locker are both SQLite now; Display is a kiosk snapshot."""

    def test_inventory_is_sqlite_catalog(
        self, client, test_devices, catalog_device
    ):
        """Inventory returns every catalog row, locker and non-locker alike."""
        resp = client.get("/api/dashboard/inventory")
        assert resp.status_code == 200
        rows = resp.json()
        pms = {r["pm_number"] for r in rows}
        assert "PM-001" in pms
        assert "PM-999" in pms
        van = next(r for r in rows if r["pm_number"] == "PM-999")
        assert van["name"] == "Van kit"
        assert van["location"] == "Workshop"
        assert van["in_locker"] is False
        locker = next(r for r in rows if r["pm_number"] == "PM-001")
        assert locker["in_locker"] is True
        assert locker["location"] == "Locker"
        assert "locker_slot" not in van
        assert "tag_hmac" not in van
        assert "uid_hmac" not in van

    def test_locker_stays_registered_only(
        self, client, test_devices, catalog_device
    ):
        """Locker tab JSON is cabinet units; catalog-only rows are absent."""
        resp = client.get("/api/dashboard/devices")
        assert resp.status_code == 200
        pms = {r["pm_number"] for r in resp.json()}
        assert "PM-001" in pms
        assert "PM-999" not in pms
        row = next(r for r in resp.json() if r["pm_number"] == "PM-001")
        assert "locker_slot" in row
        assert "status" in row

    def test_inventory_mirror_down_does_not_break_reads(
        self, client, test_devices, tmp_path, monkeypatch
    ):
        """A missing mirror file does not touch SQLite-backed reads."""
        monkeypatch.setenv(
            "SMART_LOCKER_MIRROR_PATH", str(tmp_path / "missing.xlsx")
        )
        inv = client.get("/api/dashboard/inventory")
        assert inv.status_code == 200
        locker = client.get("/api/dashboard/devices")
        assert locker.status_code == 200
        assert locker.json()

    def test_inventory_unconfigured_mirror_still_answers(self, client, test_devices):
        """No mirror configured is not an Inventory error — SQLite answers."""
        resp = client.get("/api/dashboard/inventory")
        assert resp.status_code == 200
        assert len(resp.json()) == len(test_devices)

    def test_display_idle_has_no_user(self, client, mock_context):
        """Idle kiosk reports Idle occupancy without a person name."""
        mock_context.kiosk_screen = "idle"
        mock_context.admin_overlay_open = False
        resp = client.get("/api/dashboard/display")
        assert resp.status_code == 200
        data = resp.json()
        assert data["screen"] == "idle"
        assert data["label"] == "Idle"
        assert data["occupied"] is False
        assert "user_name" not in data

    def test_display_shows_occupied_not_user_name(
        self, client, mock_context, test_user
    ):
        """Logged-in kiosk reports occupancy, not the borrower/user name."""
        mock_context.session_mgr.start_session(test_user)
        mock_context.kiosk_screen = "main-menu"
        resp = client.get("/api/dashboard/display")
        assert resp.status_code == 200
        data = resp.json()
        assert data["screen"] == "main-menu"
        assert data["label"] == "Main menu"
        assert data["occupied"] is True
        assert "user_name" not in data
        assert "Test User" not in resp.text

    def test_display_from_lan_has_no_person_name(
        self, lan_client, mock_context, test_user
    ):
        """Public Display from LAN stays occupancy-only even with a live session."""
        mock_context.session_mgr.start_session(test_user)
        resp = lan_client.get("/api/dashboard/display")
        assert resp.status_code == 200
        data = resp.json()
        assert data["occupied"] is True
        assert "user_name" not in data

    def test_heartbeat_updates_display(self, client, mock_context):
        """Kiosk POST of a screen id is what Display polls."""
        resp = client.post("/api/kiosk/display", json={"screen": "borrow"})
        assert resp.status_code == 200
        assert mock_context.kiosk_screen == "borrow"
        shown = client.get("/api/dashboard/display").json()
        assert shown["screen"] == "borrow"
        assert shown["label"] == "Locker"

    def test_display_admin_overlay_overrides_heartbeat(self, client, mock_context):
        """Open admin overlay is Admin even if the last heartbeat was idle."""
        mock_context.kiosk_screen = "idle"
        mock_context.admin_overlay_open = True
        data = client.get("/api/dashboard/display").json()
        assert data["screen"] == "admin"
        assert data["label"] == "Admin"

    def test_heartbeat_from_lan_is_403(self, lan_client, mock_context):
        """Display heartbeat is kiosk-loopback only."""
        mock_context.kiosk_screen = "idle"
        resp = lan_client.post("/api/kiosk/display", json={"screen": "borrow"})
        assert resp.status_code == 403
        assert mock_context.kiosk_screen == "idle"

    def test_heartbeat_null_context_is_ok(self, client, monkeypatch):
        """Missing AppContext on POST is 200, not 500."""
        monkeypatch.setattr(ctx_module, "context", None)
        resp = client.post("/api/kiosk/display", json={"screen": "idle"})
        assert resp.status_code == 200
        assert resp.json()["ok"] is True


class TestDashboardOwnerEditApi:
    """POST /api/dashboard/owner is public — holder of a non-cabinet row."""

    def test_owner_change_is_public(
        self, client, catalog_device, db_session
    ):
        """No secret, no session: a non-cabinet row's owner just updates."""
        resp = client.post(
            "/api/dashboard/owner",
            json={"pm_number": "PM-999", "owner": "Alex"},
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["ok"] is True
        assert body["locker"] is False
        db_session.expire_all()
        assert catalog_device.location == "Alex"

    def test_owner_change_from_lan_is_public(
        self, lan_client, catalog_device, db_session
    ):
        """Owner edit stays open for LAN staff — that is the plan's choice."""
        resp = lan_client.post(
            "/api/dashboard/owner",
            json={"pm_number": "PM-999", "owner": "Workshop"},
        )
        assert resp.status_code == 200
        db_session.expire_all()
        assert catalog_device.location == "Workshop"

    def test_locker_pm_is_refused(
        self, client, test_user, test_devices, db_session,
    ):
        """Locker PMs cannot have owner changed from the dashboard."""
        from smart_locker.database.models import TransactionLog
        from sqlalchemy import select

        db_session.commit()
        resp = client.post(
            "/api/dashboard/owner",
            json={"pm_number": "PM-001", "owner": "Test User"},
        )
        assert resp.status_code == 409
        db_session.expire_all()
        device = DeviceRepository.find_by_pm(db_session, "PM-001")
        assert device.status == DeviceStatus.AVAILABLE
        assert device.current_borrower_id is None
        logs = db_session.execute(select(TransactionLog)).scalars().all()
        assert logs == []

    def test_owners_list_is_public(
        self, client, test_user, db_session
    ):
        """GET /api/dashboard/owners is public — the owner dialog needs it."""
        RegistrantRepository.add_names(db_session, {"Bob Field"})
        db_session.commit()
        resp = client.get("/api/dashboard/owners")
        assert resp.status_code == 200
        names = resp.json()["names"]
        assert "Test User" in names
        assert "Bob Field" in names
        assert "in_locker_token" in resp.json()

    def test_users_and_transactions_require_secret(self, client):
        """Users and last-500 transactions are not public GETs."""
        assert client.get("/api/dashboard/users").status_code == 401
        assert client.get("/api/dashboard/transactions").status_code == 401

    def test_users_and_transactions_with_secret(
        self, client, test_user, dashboard_secret
    ):
        """Authorized GET returns users without hmac/uid fields."""
        users = client.get(
            "/api/dashboard/users",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert users.status_code == 200
        row = next(u for u in users.json() if u["display_name"] == "Test User")
        assert "uid_hmac" not in row
        tx = client.get(
            "/api/dashboard/transactions",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert tx.status_code == 200
        assert isinstance(tx.json(), list)

    def test_mirror_down_owner_edit_still_writes(
        self, client, catalog_device, db_session, tmp_path, monkeypatch
    ):
        """An unavailable mirror never blocks the SQLite-backed owner edit."""
        monkeypatch.setenv(
            "SMART_LOCKER_MIRROR_PATH", str(tmp_path / "missing.xlsx")
        )
        resp = client.post(
            "/api/dashboard/owner",
            json={"pm_number": "PM-999", "owner": "Alex"},
        )
        assert resp.status_code == 200
        db_session.expire_all()
        assert catalog_device.location == "Alex"


class TestDashboardTagApi:
    """Dashboard bind/unbind require the admin secret; overlay is not auth."""

    def test_bind_tag_without_secret_is_401(self, client, mock_context, test_devices):
        """POST /api/dashboard/bind-tag is 401 when the secret is unset."""
        mock_context.pending_tag_bind = None
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
        )
        assert resp.status_code == 401
        assert mock_context.pending_tag_bind is None

    def test_unbind_tag_without_secret_is_401(
        self, client, test_devices, db_session, hmac_key
    ):
        """POST /api/dashboard/unbind-tag is 401 when the secret is unset."""
        DeviceRepository.bind_tag(
            db_session,
            test_devices[0],
            compute_uid_hmac("AABBCCDD", hmac_key),
        )
        db_session.commit()
        resp = client.post(
            "/api/dashboard/unbind-tag",
            json={"pm_number": "PM-001"},
        )
        assert resp.status_code == 401
        db_session.expire_all()
        assert test_devices[0].tag_hmac is not None

    def test_bind_tag_wrong_secret_is_401(
        self, client, mock_context, test_devices, dashboard_secret
    ):
        """Wrong header value is 401, not a successful arm."""
        mock_context.pending_tag_bind = None
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers("wrong-secret-value"),
        )
        assert resp.status_code == 401
        assert mock_context.pending_tag_bind is None

    def test_bind_tag_with_secret_arms(
        self, client, mock_context, test_devices, dashboard_secret
    ):
        """Configured secret arms pending_tag_bind without a kiosk session."""
        mock_context.pending_tag_bind = None
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert resp.json().get("ok") is True
        assert mock_context.pending_tag_bind is not None
        assert mock_context.pending_tag_bind.device_id == test_devices[0].id
        assert mock_context.pending_tag_bind.is_expired is False
        assert mock_context.pending_tag_bind.from_dashboard is True

    def test_dashboard_bind_from_lan_is_secret_not_loopback(
        self, lan_client, mock_context, test_devices, dashboard_secret
    ):
        """Dashboard bind stays secret-gated for LAN staff, not loopback-only."""
        mock_context.pending_tag_bind = None
        resp = lan_client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert mock_context.pending_tag_bind is not None
        assert mock_context.pending_tag_bind.from_dashboard is True

    def test_unbind_tag_with_secret_clears_hmac(
        self, client, test_devices, db_session, hmac_key, dashboard_secret
    ):
        """Authorized unbind clears tag_hmac; has_tag becomes False."""
        DeviceRepository.bind_tag(
            db_session,
            test_devices[0],
            compute_uid_hmac("AABBCCDD", hmac_key),
        )
        db_session.commit()
        resp = client.post(
            "/api/dashboard/unbind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert resp.json().get("ok") is True
        db_session.expire_all()
        assert test_devices[0].tag_hmac is None
        listed = client.get("/api/dashboard/devices").json()
        cam = next(d for d in listed if d["name"] == "Camera")
        assert cam["has_tag"] is False
        assert "tag_hmac" not in cam

    def test_unbind_clears_pending_tag_bind(
        self, client, mock_context, test_devices, db_session, hmac_key, dashboard_secret
    ):
        """Dashboard unbind cancels an armed bind window."""
        DeviceRepository.bind_tag(
            db_session,
            test_devices[0],
            compute_uid_hmac("AABBCCDD", hmac_key),
        )
        db_session.commit()
        mock_context.pending_tag_bind = PendingTagBind(device_id=test_devices[0].id)
        resp = client.post(
            "/api/dashboard/unbind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert mock_context.pending_tag_bind is None

    def test_bind_refused_while_kiosk_session_active(
        self, client, mock_context, test_user, test_devices, dashboard_secret
    ):
        """Dashboard bind must not arm while someone is logged in at the kiosk."""
        mock_context.session_mgr.start_session(test_user)
        mock_context.pending_tag_bind = None
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409
        assert mock_context.pending_tag_bind is None

    def test_bind_does_not_drop_pending_registration(
        self, client, mock_context, test_devices, dashboard_secret
    ):
        """Dashboard bind must not clear an in-progress kiosk enroll."""
        mock_context.pending_registration = PendingRegistration("Someone")
        mock_context.pending_tag_bind = None
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409
        assert mock_context.pending_registration is not None
        assert mock_context.pending_registration.display_name == "Someone"
        assert mock_context.pending_tag_bind is None

    def test_second_bind_rejected_while_armed(
        self, client, mock_context, test_devices, dashboard_secret
    ):
        """Last-writer-wins is refused: a non-expired bind blocks a new arm."""
        mock_context.pending_tag_bind = PendingTagBind(device_id=test_devices[1].id)
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-001"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409
        assert mock_context.pending_tag_bind.device_id == test_devices[1].id

    def test_unknown_pm_is_404(
        self, client, mock_context, dashboard_secret
    ):
        """Bind/unbind of a PM that is not a locker device is 404."""
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-MISSING"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404
        resp = client.post(
            "/api/dashboard/unbind-tag",
            json={"pm_number": "PM-MISSING"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404
        assert mock_context.pending_tag_bind is None

    def test_non_locker_row_bind_is_404(
        self, client, mock_context, catalog_device, dashboard_secret
    ):
        """Tag bind/unbind apply to cabinet units only — catalog rows 404."""
        resp = client.post(
            "/api/dashboard/bind-tag",
            json={"pm_number": "PM-999"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404
        assert mock_context.pending_tag_bind is None
        resp = client.post(
            "/api/dashboard/unbind-tag",
            json={"pm_number": "PM-999"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404


class TestDashboardCatalogEditor:
    """Dashboard catalog CRUD — admin secret on mutations, public reads."""

    def test_add_device_requires_secret(self, client):
        """POST /api/dashboard/devices is 401 when the secret is unset."""
        resp = client.post(
            "/api/dashboard/devices",
            json={"pm_number": "PM-500", "name": "Crimpers"},
        )
        assert resp.status_code == 401

    def test_add_device_creates_catalog_row(
        self, client, db_session, dashboard_secret
    ):
        """An added device lands in SQLite, not in a slot, and is listed."""
        resp = client.post(
            "/api/dashboard/devices",
            json={
                "pm_number": "PM-500",
                "name": "Crimpers",
                "device_type": "Tool",
                "serial_number": "SN-500",
                "location": "Workshop",
            },
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        body = resp.json()
        assert body["ok"] is True
        assert body["device"]["pm_number"] == "PM-500"
        assert body["device"]["in_locker"] is False
        db_session.expire_all()
        row = DeviceRepository.find_by_pm(db_session, "PM-500")
        assert row is not None
        assert row.locker_slot is None
        inventory = client.get("/api/dashboard/inventory").json()
        assert "PM-500" in {r["pm_number"] for r in inventory}

    def test_add_duplicate_pm_is_409(
        self, client, catalog_device, dashboard_secret
    ):
        """Adding an id that is already cataloged is a conflict."""
        resp = client.post(
            "/api/dashboard/devices",
            json={"pm_number": "PM-999", "name": "Second van kit"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409

    def test_add_blank_name_is_422(self, client, dashboard_secret):
        """Name is required — an empty one is a validation failure."""
        resp = client.post(
            "/api/dashboard/devices",
            json={"pm_number": "PM-501", "name": "   "},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 422

    def test_edit_device_fields(
        self, client, catalog_device, db_session, dashboard_secret
    ):
        """PATCH updates catalog fields and the stored place."""
        resp = client.patch(
            "/api/dashboard/devices/PM-999",
            json={"name": "Kit van", "location": "Field", "manufacturer": "Fluke"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert resp.json()["ok"] is True
        db_session.expire_all()
        assert catalog_device.name == "Kit van"
        assert catalog_device.location == "Field"

    def test_edit_location_on_locker_row_is_409(
        self, client, test_devices, dashboard_secret
    ):
        """Place on a cabinet unit belongs to borrow/return — refuse."""
        resp = client.patch(
            "/api/dashboard/devices/PM-001",
            json={"location": "Workshop"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409

    def test_edit_metadata_on_locker_row_is_ok(
        self, client, test_devices, db_session, dashboard_secret
    ):
        """Name/serial edits stay allowed on cabinet units — only place is owned."""
        resp = client.patch(
            "/api/dashboard/devices/PM-001",
            json={"name": "Cam A"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        db_session.expire_all()
        assert test_devices[0].name == "Cam A"

    def test_edit_unknown_pm_is_404(self, client, dashboard_secret):
        resp = client.patch(
            "/api/dashboard/devices/PM-NOPE",
            json={"name": "x"},
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404

    def test_remove_device(
        self, client, catalog_device, db_session, dashboard_secret
    ):
        """DELETE removes a non-borrowed row from the catalog."""
        resp = client.delete(
            "/api/dashboard/devices/PM-999",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        db_session.expire_all()
        assert DeviceRepository.find_by_pm(db_session, "PM-999") is None

    def test_remove_borrowed_is_409(
        self, client, mock_context, test_user, test_devices, db_session,
        dashboard_secret,
    ):
        """A borrowed unit cannot be deleted — return it first."""
        from smart_locker.services.locker_service import LockerService

        user_session = mock_context.session_mgr.start_session(test_user)
        assert LockerService.borrow_device(db_session, user_session, test_devices[0].id)
        db_session.commit()
        resp = client.delete(
            "/api/dashboard/devices/PM-001",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409
        db_session.expire_all()
        assert DeviceRepository.find_by_pm(db_session, "PM-001") is not None

    def test_remove_device_with_history_is_409(
        self, client, mock_context, test_user, test_devices, db_session,
        dashboard_secret,
    ):
        """A returned unit keeps its audit rows — removal is refused, not 500."""
        from smart_locker.services.locker_service import LockerService

        user_session = mock_context.session_mgr.start_session(test_user)
        assert LockerService.borrow_device(db_session, user_session, test_devices[0].id)
        db_session.commit()
        LockerService.return_device(db_session, user_session, test_devices[0].id)
        db_session.commit()
        resp = client.delete(
            "/api/dashboard/devices/PM-001",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 409
        db_session.expire_all()
        assert DeviceRepository.find_by_pm(db_session, "PM-001") is not None

    def test_remove_unknown_pm_is_404(self, client, dashboard_secret):
        resp = client.delete(
            "/api/dashboard/devices/PM-NOPE",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 404

    def test_mutations_require_secret_even_when_set(self, client, dashboard_secret):
        """A wrong header value on CRUD is still 401."""
        resp = client.post(
            "/api/dashboard/devices",
            json={"pm_number": "PM-500", "name": "Crimpers"},
            headers=dashboard_admin_headers("wrong-secret"),
        )
        assert resp.status_code == 401
        resp = client.delete(
            "/api/dashboard/devices/PM-500",
            headers=dashboard_admin_headers("wrong-secret"),
        )
        assert resp.status_code == 401


class TestDashboardMirrorApi:
    """Mirror status/diffs/apply/dismiss endpoints."""

    def _configure_mirror(self, monkeypatch, tmp_path, path):
        monkeypatch.setenv("SMART_LOCKER_MIRROR_PATH", str(path))

    def test_mirror_status_public_unconfigured(self, client):
        """GET /api/dashboard/mirror is public and reports unconfigured."""
        resp = client.get("/api/dashboard/mirror")
        assert resp.status_code == 200
        body = resp.json()
        assert body["configured"] is False
        assert body["state"] == "unconfigured"

    def test_mirror_diffs_requires_secret(self, client):
        resp = client.get("/api/dashboard/mirror/diffs")
        assert resp.status_code == 401

    def test_mirror_apply_requires_secret(self, client):
        resp = client.post("/api/dashboard/mirror/apply")
        assert resp.status_code == 401

    def test_mirror_dismiss_requires_secret(self, client):
        resp = client.post("/api/dashboard/mirror/dismiss")
        assert resp.status_code == 401

    def test_mirror_apply_unavailable_is_503(
        self, client, tmp_path, monkeypatch, dashboard_secret
    ):
        """Applying with no readable mirror file maps to 503."""
        self._configure_mirror(monkeypatch, tmp_path, tmp_path / "gone.xlsx")
        resp = client.post(
            "/api/dashboard/mirror/apply",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 503

    def test_external_edit_detected_and_dismissed(
        self, client, catalog_device, db_session, tmp_path, monkeypatch,
        dashboard_secret,
    ):
        """A hand edit flips external_changes; dismiss keeps the database."""
        from smart_locker.database.engine import get_engine
        from smart_locker.sync import mirror
        from openpyxl import load_workbook

        path = tmp_path / "mirror.xlsx"
        self._configure_mirror(monkeypatch, tmp_path, path)
        engine = get_engine()
        # First tick adopts/seeds; pending write flushes the catalog out.
        mirror.tick(engine, trigger="manual")
        assert path.exists()
        # Hand-edit the file: change the catalog row's Location cell.
        wb = load_workbook(path)
        ws = wb.active
        ws.cell(row=2, column=8, value="Edited by hand")
        wb.save(path)
        mirror.tick(engine, trigger="manual")
        status = client.get("/api/dashboard/mirror").json()
        assert status["external_changes"] is True
        diffs = client.get(
            "/api/dashboard/mirror/diffs",
            headers=dashboard_admin_headers(dashboard_secret),
        ).json()["diffs"]
        assert any(d["kind"] == "changed" for d in diffs)
        resp = client.post(
            "/api/dashboard/mirror/dismiss",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        db_session.expire_all()
        assert catalog_device.location == "Workshop"
        status = client.get("/api/dashboard/mirror").json()
        assert status["external_changes"] is False

        # Dismissed edits do not stay in the file: the scheduled flush
        # rewrites the sheet from the database.
        mirror.flush_scheduled()
        mirror.tick(engine, trigger="manual")
        wb = load_workbook(path, read_only=True)
        try:
            row = list(wb.active.iter_rows(min_row=2, values_only=True))[0]
        finally:
            wb.close()
        assert row[7] == "Workshop"

    def test_apply_added_row_onto_existing_device_commits(
        self, client, db_session, tmp_path, monkeypatch, dashboard_secret
    ):
        """A sheet row whose PM exists in SQLite but not in the last write
        is an 'added' diff onto an existing row — the update must commit."""
        from smart_locker.database.engine import get_engine
        from smart_locker.sync import mirror
        from openpyxl import load_workbook

        path = tmp_path / "mirror.xlsx"
        self._configure_mirror(monkeypatch, tmp_path, path)
        engine = get_engine()

        # Baseline: empty catalog writes headers only (last_write_rows = []).
        mirror.tick(engine, trigger="manual")
        assert path.exists()

        # DB row created after the baseline (raw create marks nothing dirty)
        # — its PM is in SQLite but not in the mirror baseline.
        DeviceRepository.create(
            db_session,
            name="Old name",
            device_type="Tool",
            pm_number="PM-555",
            manufacturer="Old Mfr",
            model="X1",
        )
        db_session.commit()

        # Hand-add the same PM to the sheet with different fields.
        wb = load_workbook(path)
        ws = wb.active
        ws.append(["PM-555", "Sheet name", "Tool", "Sheet Mfr",
                   "X1", "", "", "Bench"])
        wb.save(path)
        wb.close()

        mirror.tick(engine, trigger="manual")
        diffs = client.get(
            "/api/dashboard/mirror/diffs",
            headers=dashboard_admin_headers(dashboard_secret),
        ).json()["diffs"]
        assert any(d["kind"] == "added" for d in diffs)

        resp = client.post(
            "/api/dashboard/mirror/apply",
            headers=dashboard_admin_headers(dashboard_secret),
        )
        assert resp.status_code == 200
        assert resp.json()["applied"] >= 1

        db_session.expire_all()
        device = DeviceRepository.find_by_pm(db_session, "PM-555")
        assert device.name == "Sheet name"
        assert device.manufacturer == "Sheet Mfr"
        assert device.location == "Bench"

