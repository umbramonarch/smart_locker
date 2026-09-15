"""
File: test_owner_edit.py
Description: Tests for public dashboard owner edit. Inventory tab only.
             Non-locker PMs write Excel (no SQLite insert). Locker PMs are
             refused — owner is set at the kiosk, not the dashboard.
Project: smart_locker/tests
Notes: Run with: python -m pytest tests/test_owner_edit.py -v
       Uses temporary workbooks; no NFC hardware.
"""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from smart_locker.database.models import DeviceStatus, TransactionLog
from smart_locker.database.repositories import (
    DeviceRepository,
    RegistrantRepository,
    UserRepository,
)
from smart_locker.security.encryption import encrypt
from smart_locker.security.hashing import compute_uid_hmac
from smart_locker.services.owner_edit import LockerOwned, owner_choices, set_owner
from smart_locker.sync.location_writeback import IN_LOCKER_TOKEN


def _workbook(path: Path, rows: list[list]) -> Path:
    """Write a workbook (first row = headers) and return ``path``."""
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _location_by_pm(path: Path) -> dict[str, str]:
    """Read Equipment → Location from the first sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(h).strip() if h else "" for h in rows[0]]
    pm_i = headers.index("Equipment")
    loc_i = headers.index("Location")
    out: dict[str, str] = {}
    for row in rows[1:]:
        pm = str(row[pm_i]).strip() if row[pm_i] is not None else ""
        loc = "" if row[loc_i] is None else str(row[loc_i]).strip()
        if pm:
            out[pm] = loc
    return out


def _add_user(db_session, enc_key, hmac_key, name: str = "Alice"):
    """Insert one enrolled user and return it."""
    uid = f"UID-{name}"
    return UserRepository.create(
        db_session,
        display_name=name,
        uid_hmac=compute_uid_hmac(uid, hmac_key),
        encrypted_card_uid=encrypt(uid, enc_key),
    )


class TestOwnerChoices:
    """Dropdown is registered users + registrant names (+ in-locker token)."""

    def test_includes_users_registrants_and_in_locker_token(
        self, db_session, enc_key, hmac_key
    ):
        """Choices list enrolled names, Excel registrant names, and Locker."""
        _add_user(db_session, enc_key, hmac_key, "Alice")
        RegistrantRepository.add_names(db_session, {"Bob Field"})
        db_session.flush()

        names = owner_choices(db_session)

        assert IN_LOCKER_TOKEN in names
        assert "Alice" in names
        assert "Bob Field" in names

    def test_inactive_users_are_excluded(self, db_session, enc_key, hmac_key):
        """Deactivated users disappear from the owner dropdown."""
        _add_user(db_session, enc_key, hmac_key, "Alice")
        gone = _add_user(db_session, enc_key, hmac_key, "Gone")
        gone.is_active = False
        db_session.flush()

        names = owner_choices(db_session)

        assert "Alice" in names
        assert "Gone" not in names

    def test_owner_choices_exclude_deactivated_name(
        self, db_session, enc_key, hmac_key
    ):
        """A registrant name matching a deactivated user is not offered."""
        RegistrantRepository.add_names(db_session, {"Alice"})
        alice = _add_user(db_session, enc_key, hmac_key, "Alice")
        alice.is_active = False
        _add_user(db_session, enc_key, hmac_key, "Bob")
        db_session.flush()

        names = owner_choices(db_session)

        assert "Alice" not in names
        assert "Bob" in names


class TestSetOwnerExcelOnly:
    """Inventory owner edit on a non-locker PM writes Excel, not SQLite."""

    def test_non_locker_writes_excel_only(
        self, db_session, tmp_path, enc_key, hmac_key
    ):
        """Changing owner of an Excel-only PM does not create a locker row."""
        _add_user(db_session, enc_key, hmac_key, "Alice")
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Name", "Location"],
            ["PM-VAN", "Van kit", "Workshop"],
            ["PM-LOCKER", "Scope", "Locker"],
        ])
        DeviceRepository.create(
            db_session,
            name="Scope",
            device_type="general",
            pm_number="PM-LOCKER",
            locker_slot=1,
        )
        db_session.flush()

        result = set_owner(db_session, path, "PM-VAN", "Alice")

        assert result.locker is False
        assert _location_by_pm(path)["PM-VAN"] == "Alice"
        assert _location_by_pm(path)["PM-LOCKER"] == "Locker"
        assert DeviceRepository.find_by_pm(db_session, "PM-VAN") is None
        assert db_session.execute(select(TransactionLog)).scalars().all() == []


class TestSetOwnerLockerRefused:
    """Locker devices stay kiosk-owned. Dashboard must not change them."""

    def test_locker_pm_is_refused(
        self, db_session, tmp_path, enc_key, hmac_key
    ):
        """A locker PM is not written in Excel or SQLite from the dashboard."""
        _add_user(db_session, enc_key, hmac_key, "Alice")
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Name", "Location"],
            ["PM-001", "Scope", "Locker"],
        ])
        device = DeviceRepository.create(
            db_session,
            name="Scope",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
            status=DeviceStatus.AVAILABLE.value,
        )
        db_session.flush()

        with pytest.raises(LockerOwned):
            set_owner(db_session, path, "PM-001", "Alice")

        db_session.refresh(device)
        assert _location_by_pm(path)["PM-001"] == "Locker"
        assert device.status == DeviceStatus.AVAILABLE
        assert device.current_borrower_id is None
        assert db_session.execute(select(TransactionLog)).scalars().all() == []
