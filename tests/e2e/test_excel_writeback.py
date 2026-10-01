"""
File: test_excel_writeback.py
Description: End-to-end coverage of the physical Excel Location write-back —
             the one sync path the API suite always mocks. Boots create_app()
             via the ``e2e`` factory against a real catalog workbook on disk:
             card/sticker taps drive borrow/return through the NFC bridge and
             LockerService, then flush_scheduled_writeback() drains the async
             writer thread before the xlsx is reopened and the Location cell
             is asserted with the writer's own column auto-detection.
Project: smart_locker/tests/e2e
Notes: SOURCE_EXCEL_PATH is monkeypatched per test — the suite autouse fixture
       pins it to "" so a developer .env can never redirect the write-back.
       Workbooks need a PM header ("Equipment" is a pm_candidate) and a
       Location header ("Location" is a location_candidate) or the writer
       skips the sheet with an error result.
"""

from openpyxl import load_workbook

from config.settings import in_locker_token
from smart_locker.sync.location_writeback import flush_scheduled_writeback
from smart_locker.sync.source_import import (
    find_column,
    location_candidates,
    pm_candidates,
    pm_match_key,
)
from tests.api.helpers import catalog_workbook
from tests.e2e.helpers import add_device, add_user, get_device

# Simulated UIDs — hex-like strings. The DB is fresh per test, so constants
# can repeat safely across tests.
CARD_USER = "0A30000001"
CARD_ADMIN = "0A30000002"
TAG_PM100 = "0A40000001"
TAG_PM999 = "0A40000002"

# Header row used by every workbook in this file — "Equipment" and "Location"
# are the first PM/Location candidates, so detection is deterministic.
CATALOG_HEADERS = ["Equipment", "Name", "Type", "Location"]


def _login(h, card_uid: str) -> dict:
    """Tap a work card and wait for the session to open."""
    h.tap(card_uid)
    return h.wait_event("auth_success")


def _point_source_at(monkeypatch, path) -> None:
    """Aim the write-back at this test's workbook (override the autouse "")."""
    monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))


def _location_cell(path, pm_number: str):
    """Reopen the workbook and return the Location cell text for one PM.

    Uses the same header auto-detection and PM join key as the writer, so the
    row is found however the sheet spells its headers.

    Args:
        path: Path to the device-list.xlsx under test.
        pm_number: Equipment number whose Location cell is read.

    Returns:
        The cell value (usually str), or raises AssertionError when the PM
        row or the PM/Location columns are absent.
    """
    wb = load_workbook(path)
    try:
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        pm_idx = find_column(headers, pm_candidates())
        loc_idx = find_column(headers, location_candidates())
        assert pm_idx is not None, f"No PM column in {path} (headers: {headers})"
        assert loc_idx is not None, (
            f"No Location column in {path} (headers: {headers})"
        )
        want = pm_match_key(pm_number)
        for row_i in range(2, ws.max_row + 1):
            pm = ws.cell(row=row_i, column=pm_idx + 1).value
            if pm_match_key(pm) == want:
                return ws.cell(row=row_i, column=loc_idx + 1).value
        raise AssertionError(f"PM {pm_number} not found in {path}")
    finally:
        wb.close()


def test_borrow_writes_borrower_name_to_location(e2e, tmp_path, monkeypatch):
    """Sticker tap borrow lands the borrower's display name in Excel Location."""
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "Locker"],
    ])
    _point_source_at(monkeypatch, path)
    h = e2e()
    user_id = add_user(h, CARD_USER, display_name="E2E Borrower")
    device_id = add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)

    payload = h.wait_event("device_action")
    assert payload["success"] is True
    assert payload["action"] == "borrow"
    assert payload["device_id"] == device_id

    device = get_device(h, device_id)
    assert device.current_borrower_id == user_id

    # The write-back runs on a worker thread — drain it before reading cells.
    flush_scheduled_writeback(timeout=10)
    assert _location_cell(path, "PM-100") == "E2E Borrower"


def test_return_writes_in_locker_token(e2e, tmp_path, monkeypatch):
    """A return puts the in-locker token (not a blank) back into Location."""
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "Locker"],
    ])
    _point_source_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)
    borrow = h.wait_event("device_action")
    assert borrow["action"] == "borrow"
    flush_scheduled_writeback(timeout=10)
    # Borrower name is on the sheet while the device is out.
    assert _location_cell(path, "PM-100") == "E2E Borrower"

    h.tap(TAG_PM100)
    returned = h.wait_event("device_action")
    assert returned["success"] is True
    assert returned["action"] == "return"

    flush_scheduled_writeback(timeout=10)
    assert _location_cell(path, "PM-100") == in_locker_token()


def test_maintenance_device_location_not_written(e2e, tmp_path, monkeypatch):
    """MAINTENANCE rows are excluded from the wanted map — the writer walks
    their Excel row but never touches the Location cell."""
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "Locker"],
        ["PM-999", "In Repair", "Tool", "Cabinet 9"],
    ])
    _point_source_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )
    add_device(
        h, name="In Repair", pm_number="PM-999", locker_slot=2,
        tag_uid=TAG_PM999, status="maintenance",
    )

    _login(h, CARD_USER)

    # A maintenance sticker cannot be borrowed — the tap is refused, so no
    # write-back is even scheduled for it.
    h.tap(TAG_PM999)
    refused = h.wait_event("device_action")
    assert refused["success"] is False
    assert refused["action"] == "refused"

    # A real borrow drives the write-back pass over the whole sheet.
    h.tap(TAG_PM100)
    borrowed = h.wait_event("device_action")
    assert borrowed["success"] is True

    flush_scheduled_writeback(timeout=10)
    assert _location_cell(path, "PM-100") == "E2E Borrower"
    # The maintenance PM row keeps whatever the sheet already had.
    assert _location_cell(path, "PM-999") == "Cabinet 9"


def test_health_reports_last_writeback(e2e, tmp_path, monkeypatch):
    """GET /api/health exposes the write-back snapshot after a flushed write."""
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "Locker"],
    ])
    _point_source_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)
    h.wait_event("device_action")
    flush_scheduled_writeback(timeout=10)

    resp = h.client.get("/api/health")
    assert resp.status_code == 200
    body = resp.json()
    wb = body["last_writeback"]
    assert wb is not None
    assert wb["error"] is None
    assert wb["saved"] is True
    assert wb["written"] >= 1
    assert wb["at"]  # ISO timestamp of the last attempt


def test_register_device_writes_in_locker_token(e2e, tmp_path, monkeypatch):
    """Admin Register Device inserts the row AND schedules the write-back —
    a fresh AVAILABLE device lands the in-locker token on its catalog row."""
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-200", "Multimeter", "Tool", "Cabinet 2"],
    ])
    _point_source_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_ADMIN, display_name="Site Admin", role="admin")

    _login(h, CARD_ADMIN)
    resp = h.client.post(
        "/api/admin/devices/register",
        json={"pm_number": "PM-200", "locker_slot": 7},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["pm_number"] == "PM-200"
    assert body["locker_slot"] == 7

    flush_scheduled_writeback(timeout=10)
    assert _location_cell(path, "PM-200") == in_locker_token()
