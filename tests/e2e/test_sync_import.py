"""
File: test_sync_import.py
Description: True end-to-end coverage of the source-Excel sync path. Real
             workbooks on disk are imported through POST /api/admin/sync-source,
             the dry-run preview, and the scheduler (startup + watchdog file
             triggers) — every write lands in the shared in-memory DB, and the
             outcome is read back through /api/admin/sync-status and /api/health.
Project: smart_locker/tests/e2e
Notes: Workbooks are built with tests.api.helpers.catalog_workbook (same helper
       as the API tests) and SOURCE_EXCEL_PATH is pointed at them per test —
       the route reads it at request time. Admin calls ride the real session
       manager via POST /api/admin/session (loopback). Scheduler tests call
       stop_scheduler() in a finally block; the app lifespan calls it too and
       it is idempotent. The 409 test holds the real scheduler._import_lock on
       the test thread so run_source_import_exclusive raises ImportInProgress
       deterministically inside the request handler.
"""

import time

from smart_locker.database.engine import get_engine
from smart_locker.database.models import DeviceStatus
from smart_locker.database.repositories import DeviceRepository, RegistrantRepository
from smart_locker.sync import scheduler, sync_status
from tests.api.helpers import catalog_workbook
from tests.e2e.helpers import add_device, add_user, get_device, uid_hmac_for

# Simulated UIDs — hex-like strings, unique within each test. The DB is fresh
# per test, so constants can repeat safely across tests.
ADMIN_CARD = "0A300000A1"
DEVICE_TAG = "0A300000D1"

# Header row the importer's column auto-detection resolves: PM join key, the
# catalog fields it may refresh, and the Location column used for registrants
# and write-back.
HEADERS = ["PM", "Name", "Type", "Manufacturer", "Model", "Serial", "Location"]


def _admin_session(h) -> dict:
    """Enroll an admin user and open the loopback admin-panel session."""
    add_user(h, ADMIN_CARD, display_name="E2E Admin", role="admin")
    resp = h.client.post("/api/admin/session")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["success"] is True
    assert body["user"]["role"] == "admin"
    return body


def _use_source(monkeypatch, path) -> None:
    """Point SOURCE_EXCEL_PATH at a temp workbook (read at request time)."""
    monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))


def _wait_for_trigger(trigger: str, timeout: float) -> dict:
    """Poll sync_status.get() until the last import ran under ``trigger``."""
    deadline = time.monotonic() + timeout
    last: dict = {}
    while time.monotonic() < deadline:
        snap = sync_status.get()
        if snap.get("at") and snap.get("trigger") == trigger:
            return snap
        last = snap
        time.sleep(0.05)
    raise TimeoutError(f"No '{trigger}' import within {timeout}s. Last: {last}")


def _location_for_pm(path, pm_number: str):
    """Read back the Location cell of one PM row from the workbook on disk."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(h).strip() if h else "" for h in rows[0]]
    pm_col, loc_col = headers.index("PM"), headers.index("Location")
    for row in rows[1:]:
        if row[pm_col] is not None and str(row[pm_col]).strip() == pm_number:
            return None if row[loc_col] is None else str(row[loc_col]).strip()
    return None


def test_sync_source_endpoint_updates_locker_rows_from_workbook(
    e2e, tmp_path, monkeypatch
):
    """Manual Sync runs a real import: catalog fields update, protected
    locker state and the tag binding survive, non-locker PMs are never
    inserted, Location names become registrants, and sync-status + health
    report the manual run — all through real HTTP + real SQLite."""
    h = e2e()
    device_id = add_device(
        h, name="Old Scope", pm_number="PM-100", locker_slot=5,
        tag_uid=DEVICE_TAG,
    )
    path = catalog_workbook(tmp_path, [
        HEADERS,
        # Locker PM: every catalog field differs -> counted as updated.
        ["PM-100", "Rigol DS1054", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-RIGOL-1", ""],
        # Non-locker PMs: skipped on insert, Location names -> registrants.
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Alice Smith"],
        ["PM-901", "Hakko FX-951", "Soldering Station", "Hakko", "FX-951",
         "SN-HAKKO-3", "Bob Jones"],
    ])
    _use_source(monkeypatch, path)
    _admin_session(h)

    resp = h.client.post("/api/admin/sync-source")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {
        "success": True,
        "imported": 0,
        "updated": 1,
        "unchanged": 0,
        "errors": 0,
    }

    # Catalog fields refreshed on the existing locker row.
    device = get_device(h, device_id)
    assert device.name == "Rigol DS1054"
    assert device.device_type == "Oscilloscope"
    assert device.manufacturer == "Rigol"
    assert device.model == "DS1054Z"
    assert device.serial_number == "SN-RIGOL-1"

    # Protected fields are never overwritten by a source import.
    assert device.status == DeviceStatus.AVAILABLE
    assert device.locker_slot == 5
    assert device.tag_hmac == uid_hmac_for(DEVICE_TAG)

    # The never-inserts rule: Excel-only PMs stay out of the locker DB.
    with h.db() as db:
        assert DeviceRepository.find_by_pm(db, "PM-900") is None
        assert DeviceRepository.find_by_pm(db, "PM-901") is None
        registrants = {r.display_name for r in RegistrantRepository.get_all(db)}
    assert registrants == {"Alice Smith", "Bob Jones"}
    # The public registration list reflects the synced names (Locker token is
    # an in-locker marker, never a registrant).
    assert h.client.get("/api/registrants").json()["names"] == [
        "Alice Smith", "Bob Jones",
    ]

    # The run is recorded for the admin footer and the health probe.
    status = h.client.get("/api/admin/sync-status").json()
    assert status["trigger"] == "manual"
    assert status["ok"] is True
    assert status["imported"] == 0
    assert status["updated"] == 1
    assert status["errors"] == 0
    assert status["at"]
    assert h.client.get("/api/health").json()["last_sync"]["trigger"] == "manual"

    # Location write-back ran in the same request chain: the locker PM's
    # empty Location cell now carries the in-locker token.
    assert _location_for_pm(path, "PM-100") == "Locker"


def test_sync_preview_reports_diff_and_writes_nothing(
    e2e, tmp_path, monkeypatch
):
    """Dry-run preview returns diff counts but commits nothing: device rows,
    the registrants table, the workbook file, and sync-status stay untouched."""
    h = e2e()
    device_id = add_device(h, name="Old Scope", pm_number="PM-100", locker_slot=5)
    path = catalog_workbook(tmp_path, [
        HEADERS,
        ["PM-100", "Rigol DS1054", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-RIGOL-1", ""],
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Alice Smith"],
    ])
    _use_source(monkeypatch, path)
    _admin_session(h)
    workbook_bytes = path.read_bytes()

    resp = h.client.post("/api/admin/sync-preview")
    assert resp.status_code == 200, resp.text
    assert resp.json() == {
        "preview": True,
        "imported": 0,
        "updated": 1,
        "unchanged": 0,
        "skipped": 1,
        "errors": 0,
    }

    device = get_device(h, device_id)
    assert device.name == "Old Scope"
    assert device.serial_number is None
    with h.db() as db:
        assert RegistrantRepository.get_all(db) == []
    assert path.read_bytes() == workbook_bytes

    snap = h.client.get("/api/admin/sync-status").json()
    assert snap["at"] is None
    assert snap["trigger"] is None


def test_sync_source_rejects_unconfigured_or_missing_path(
    e2e, tmp_path, monkeypatch
):
    """Empty SOURCE_EXCEL_PATH -> 400; a configured but missing file -> 400
    (run_source_import_exclusive returns None for a missing workbook)."""
    h = e2e()
    _admin_session(h)

    monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", "")
    resp = h.client.post("/api/admin/sync-source")
    assert resp.status_code == 400
    assert "not configured" in resp.json()["detail"].lower()

    _use_source(monkeypatch, tmp_path / "missing.xlsx")
    resp = h.client.post("/api/admin/sync-source")
    assert resp.status_code == 400
    assert "not found" in resp.json()["detail"].lower()


def test_sync_source_returns_409_when_import_lock_is_held(
    e2e, tmp_path, monkeypatch
):
    """Holding the real import mutex makes the request thread's
    run_source_import_exclusive raise ImportInProgress -> HTTP 409."""
    h = e2e()
    path = catalog_workbook(tmp_path, [
        HEADERS,
        ["PM-100", "Scope", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-1", "Locker"],
    ])
    _use_source(monkeypatch, path)
    _admin_session(h)

    scheduler._import_lock.acquire()
    try:
        resp = h.client.post("/api/admin/sync-source")
        assert resp.status_code == 409
        assert "already running" in resp.json()["detail"].lower()
    finally:
        scheduler._import_lock.release()


def test_start_scheduler_runs_immediate_startup_import(e2e, tmp_path):
    """start_scheduler performs one synchronous import (trigger=startup)
    before its interval job and file watcher are armed."""
    h = e2e()
    device_id = add_device(h, name="Old Name", pm_number="PM-100", locker_slot=3)
    path = catalog_workbook(tmp_path, [
        HEADERS,
        ["PM-100", "Boot Scope", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-BOOT", "Locker"],
    ])

    try:
        scheduler.start_scheduler(get_engine(), str(path), interval_minutes=1)
        snap = _wait_for_trigger("startup", timeout=5.0)
        assert snap["ok"] is True
        assert snap["updated"] == 1
        assert get_device(h, device_id).name == "Boot Scope"
        assert get_device(h, device_id).serial_number == "SN-BOOT"
    finally:
        scheduler.stop_scheduler()


def test_cifs_scheduler_polls_every_five_minutes_by_default(tmp_path, monkeypatch):
    """CIFS skips its unreliable watcher and receives the five-minute interval."""
    jobs = []

    class FakeScheduler:
        def add_job(self, *args, **kwargs):
            jobs.append((args, kwargs))

        def start(self):
            pass

        def shutdown(self, wait=False):
            pass

    source = tmp_path / "device-list.xlsx"
    source.touch()
    monkeypatch.setattr(scheduler, "BackgroundScheduler", FakeScheduler)
    monkeypatch.setattr(scheduler, "is_network_path", lambda path: True)
    monkeypatch.setattr(scheduler, "run_source_import_exclusive", lambda *args, **kwargs: None)
    monkeypatch.setattr(scheduler, "_try_dashboard_launcher", lambda: None)

    try:
        scheduler.start_scheduler(object(), source)
        assert scheduler._observer is None
        assert len(jobs) == 1
        assert jobs[0][1]["trigger"].interval.total_seconds() == 5 * 60
    finally:
        scheduler.stop_scheduler()


def test_file_watch_reimports_when_workbook_changes(e2e, tmp_path):
    """Rewriting the source workbook fires the debounced watchdog import
    (trigger=watch) and the new catalog fields land in the locker DB."""
    h = e2e()
    device_id = add_device(h, name="Watch V1", pm_number="PM-100", locker_slot=3)
    path = catalog_workbook(tmp_path, [
        HEADERS,
        ["PM-100", "Watch V1", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-W1", "Locker"],
    ])

    try:
        scheduler.start_scheduler(get_engine(), str(path), interval_minutes=1)
        assert _wait_for_trigger("startup", timeout=5.0)["ok"] is True

        # Let the watchdog observer arm, then rewrite the file. Location stays
        # "Locker" so write-back finds nothing to save and no further events
        # keep the debounce alive after the watch import.
        time.sleep(0.5)
        catalog_workbook(tmp_path, [
            HEADERS,
            ["PM-100", "Watch V2", "Oscilloscope", "Rigol", "DS1054Z",
             "SN-W1", "Locker"],
        ])

        snap = _wait_for_trigger("watch", timeout=8.0)
        assert snap["ok"] is True
        assert get_device(h, device_id).name == "Watch V2"
    finally:
        scheduler.stop_scheduler()
