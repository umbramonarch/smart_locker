"""
File: test_mirror_hardening.py
Description: Adversarial-review hardening coverage for the catalog mirror.
             A missing share file must not burn the one-time adoption; an
             absent/unmounted share directory must defer creation instead of
             planting a shadow file; quarantined hand edits survive stat
             failures; a poisoned mirror_state.json cannot raise through
             mark_dirty/dismiss_external; a same-mtime size-changing edit is
             still detected; a timed-out write that lands late is not flagged
             as a hand edit; a busy reader fails fast; write_sheet rewrites
             the PM-column sheet rather than whatever sheet is active; and
             apply_external skips a diff whose commit raises IntegrityError
             or whose field write flushes into an OperationalError, while a
             unit registered after the baseline write diffs as "added" yet
             still takes the maintenance word.
Project: smart_locker/tests/e2e
Notes: start_scheduler is monkeypatched off per test so the only ticks are
       the ones the test drives — adopt/detect/flush ordering stays
       deterministic. The mirror path points at tmp_path per test.
"""

import json
import os
import threading
import time

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session

from config.settings import in_locker_token, maintenance_token
from smart_locker.database.engine import get_engine
from smart_locker.database.models import DeviceStatus
from smart_locker.database.repositories import DeviceRepository, RegistrantRepository
from smart_locker.services.device_catalog import registerable_devices
from smart_locker.sync import mirror, scheduler, sync_status
from smart_locker.sync.catalog_sheet import MIRROR_HEADERS
from smart_locker.sync.workbook_adapter import WorkbookAdapter
from tests.api.helpers import catalog_workbook, dashboard_admin_headers
from tests.e2e.helpers import add_device, add_user
from tests.e2e.test_mirror import (
    CATALOG_HEADERS,
    _drain_mirror,
    _mirror_cell,
    _point_mirror_at,
    _tick_now,
)


def _no_scheduler(monkeypatch) -> None:
    """Disable the lifespan scheduler so ticks only run when the test calls them."""
    monkeypatch.setattr(scheduler, "start_scheduler", lambda *a, **k: None)


def _raising_stat(exc: Exception):
    """A ``_stat_path`` stand-in that raises ``exc`` for any path."""

    def _stat(_path):
        raise exc

    return _stat


def _write_sheet(path, rows) -> None:
    """Write a workbook directly at ``path`` (a real sheet appearing later)."""
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def test_missing_file_creates_then_foreign_sheet_is_reviewed(
    e2e, tmp_path, monkeypatch
):
    """The flush creating a missing mirror file IS the seed — the written
    baseline is trusted from that moment. A foreign sheet dropped over it
    afterwards is a hand edit: flagged for review, never silently adopted.
    (A review-cycle regression test: seeded=False here would let the next
    tick adopt the foreign content unreviewed.)"""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    # Missing file: created by the flush — that Pi write seeds the baseline.
    result = _tick_now()
    assert result["error"] is None
    assert path.exists()
    assert _mirror_cell(path, "PM-100") == in_locker_token()
    state = mirror._load_state()
    assert state["seeded"] is True
    assert state["external_pending"] is False

    # A foreign catalog sheet lands at the path — held for review.
    _write_sheet(path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Alice Smith"],
    ])
    result = _tick_now()
    assert result["external"] is True
    state = mirror._load_state()
    assert state["external_pending"] is True

    # Admin applies: the row lands through the review gate, not silently.
    applied = mirror.apply_external(get_engine())
    assert applied["applied"] >= 1
    with h.db() as db:
        adopted = DeviceRepository.find_by_pm(db, "PM-900")
        assert adopted is not None
        assert adopted.location == "Alice Smith"


def test_unmounted_share_defers_then_real_sheet_is_adopted(
    e2e, tmp_path, monkeypatch
):
    """Boot-before-mount lifecycle: the share directory is absent so the
    pending write defers (no shadow file, no adoption consumed); when the
    mount lands with a real catalog sheet it is adopted on first sight."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mnt_locker" / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    # "Unmounted": parent dir missing → deferred, seeded NOT consumed.
    result = _tick_now()
    assert result["error"] == "unavailable"
    assert not path.exists()
    assert mirror._load_state()["seeded"] is False

    # The mount lands carrying the real catalog → adopted on first sight.
    path.parent.mkdir(parents=True)
    _write_sheet(path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Alice Smith"],
    ])
    result = _tick_now()
    assert result["error"] is None
    assert result["seeded"] is True
    assert result["adopted"] >= 1

    state = mirror._load_state()
    assert state["seeded"] is True
    assert state["external_pending"] is False
    with h.db() as db:
        adopted = DeviceRepository.find_by_pm(db, "PM-900")
        assert adopted is not None
        assert adopted.location == "Alice Smith"
        registrants = {r.display_name for r in RegistrantRepository.get_all(db)}
    assert "Alice Smith" in registrants
    # The adopted sheet is regenerated canonically on the same tick.
    assert _mirror_cell(path, "PM-900") == "Alice Smith"


def test_flush_refuses_create_into_missing_directory(
    e2e, tmp_path, monkeypatch
):
    """A pending write into a missing directory must defer — no mkdir, no
    shadow file (boot-before-mount on the share)."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "not_mounted" / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] == "unavailable"
    assert not path.exists()
    assert not path.parent.exists()  # no mkdir side effect
    state = mirror._load_state()
    assert state["seeded"] is False
    assert state["pending_writes"] is True  # deferred, not dropped
    assert state["last_error"] == "unavailable"
    # Public last-sync snapshot carries the category, not the raw path.
    assert sync_status.get()["message"] == "unavailable"


def test_adoption_place_tokens_and_registrant_exclusions(
    e2e, tmp_path, monkeypatch
):
    """The Maintenance token and in-locker token never seed registrants; a
    whole-word locker place ("Cabinet", "locker room") stores the in-locker
    token so the adopted row is registerable instead of stranded in limbo —
    and neither spelling is a person name for the registrant list."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-1", "Alpha", "Tool", "M", "Mod", "S1", "Maintenance"],
        ["PM-2", "Beta", "Tool", "M", "Mod", "S2", "Cabinet"],
        ["PM-3", "Gamma", "Tool", "M", "Mod", "S3", "Locker"],
        ["PM-4", "Delta", "Tool", "M", "Mod", "S4", "Alice Smith"],
        ["PM-5", "Epsilon", "Tool", "M", "Mod", "S5", "locker room"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["error"] is None
    assert result["seeded"] is True

    with h.db() as db:
        registrants = {r.display_name for r in RegistrantRepository.get_all(db)}
        assert maintenance_token() not in registrants
        assert in_locker_token() not in registrants
        assert "Maintenance" not in registrants
        # Whole-word locker spellings are places, not person names.
        assert "Cabinet" not in registrants
        assert "locker room" not in registrants
        assert "Alice Smith" in registrants

        pm1 = DeviceRepository.find_by_pm(db, "PM-1")
        pm2 = DeviceRepository.find_by_pm(db, "PM-2")
        pm3 = DeviceRepository.find_by_pm(db, "PM-3")
        pm4 = DeviceRepository.find_by_pm(db, "PM-4")
        pm5 = DeviceRepository.find_by_pm(db, "PM-5")
        assert pm1.location == maintenance_token()
        assert pm2.location == in_locker_token()   # "Cabinet" → token, not limbo
        assert pm3.location == in_locker_token()
        assert pm4.location == "Alice Smith"
        assert pm5.location == in_locker_token()   # "locker room" → token

        registerable = {d.pm_number for d in registerable_devices(db)}
    assert {"PM-2", "PM-3", "PM-5"} <= registerable
    assert "PM-1" not in registerable

    # The regenerated sheet writes the canonical token for in-locker rows.
    assert _mirror_cell(path, "PM-2") == in_locker_token()

    # A follow-up tick sees its own regenerated sheet — no self-quarantine.
    result = _tick_now()
    assert result["external"] is False
    assert mirror._load_state()["external_pending"] is False


def test_stat_failure_preserves_external_quarantine(e2e, tmp_path, monkeypatch):
    """A stat failure that is not FileNotFoundError (dead share, wedged
    reader, I/O timeout) must keep external_pending and never let a pending
    write flush over the held edits."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V", "SN-9", "Workshop"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["seeded"] is True
    assert _mirror_cell(path, "PM-900") == "Workshop"

    # Hand-edit the file and get it quarantined.
    wb = load_workbook(path)
    ws = wb.active
    ws.cell(row=2, column=MIRROR_HEADERS.index("Location") + 1,
            value="Moved by hand")
    wb.save(path)
    wb.close()
    result = _tick_now()
    assert result["external"] is True
    seen_mtime = mirror._load_state()["last_seen_mtime"]
    assert seen_mtime is not None

    # Now the share dies mid-review: stat raises a generic OSError. The flag
    # and the last-seen baseline must survive, and no write may land.
    mirror.mark_dirty()
    monkeypatch.setattr(
        mirror, "_stat_path", _raising_stat(OSError("simulated dead share"))
    )
    result = _tick_now()
    assert result["error"] == "unavailable"
    assert result["skipped"] == "external_changes"
    assert result["flushed"] is False
    state = mirror._load_state()
    assert state["external_pending"] is True
    assert state["last_seen_mtime"] == seen_mtime
    assert state["last_error"] == "unavailable"
    assert _mirror_cell(path, "PM-900") == "Moved by hand"

    # A timeout-style stat failure preserves the quarantine identically.
    monkeypatch.setattr(
        mirror, "_stat_path", _raising_stat(TimeoutError("workbook I/O timeout"))
    )
    result = _tick_now()
    assert result["error"] == "timeout"
    assert mirror._load_state()["external_pending"] is True
    assert _mirror_cell(path, "PM-900") == "Moved by hand"
    with h.db() as db:
        assert DeviceRepository.find_by_pm(db, "PM-900").location == "Workshop"


def test_poisoned_state_file_never_raises(e2e, tmp_path, monkeypatch):
    """A hand-edited mirror_state.json drops to per-key defaults; saves and
    merges coerce comparisons so mark_dirty/dismiss_external cannot raise
    into a committed borrow."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    mirror._state_path().write_text(json.dumps({
        "seeded": "yes",
        "pending_writes": 5,
        "external_pending": "maybe",
        "external_decided_at": 5,
        "last_write_at": 12345,
        "last_write_rows": "not-a-list",
        "last_seen_mtime": "yesterday",
        "last_seen_size": "big",
        "last_error": 42,
    }), encoding="utf-8")

    mirror.mark_dirty()  # would TypeError on the 5 > "" merge before the fix
    state = mirror._load_state()
    assert state["seeded"] is False            # "yes" dropped to default
    assert state["external_decided_at"] is None
    assert state["pending_writes"] is True     # mark_dirty landed
    assert state["last_write_rows"] == []
    assert state["last_seen_mtime"] is None
    assert state["last_seen_size"] is None
    assert state["last_error"] is None

    # A save whose in-memory decided_at is a string against the poisoned
    # disk value exercises the coerced merge comparison.
    mirror.dismiss_external()
    mirror.flush_scheduled()
    state = mirror._load_state()
    assert state["external_pending"] is False
    assert isinstance(state["external_decided_at"], str)

    # And the mirror still works end to end.
    result = _tick_now()
    assert result["error"] is None
    assert path.exists()
    assert _mirror_cell(path, "PM-100") == in_locker_token()


def test_same_mtime_size_change_is_flagged(e2e, tmp_path, monkeypatch):
    """An edit that changes file size but keeps the same mtime (filesystem
    timestamp granularity) must still be read and flagged."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V", "SN-9", "Bench"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["seeded"] is True
    assert _mirror_cell(path, "PM-900") == "Bench"

    st = path.stat()
    wb = load_workbook(path)
    ws = wb.active
    ws.cell(row=2, column=MIRROR_HEADERS.index("Location") + 1,
            value="A much longer hand-typed location")
    wb.save(path)
    wb.close()
    assert path.stat().st_size != st.st_size
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns))  # same mtime as baseline

    result = _tick_now()
    assert result["external"] is True
    assert mirror._load_state()["external_pending"] is True
    with h.db() as db:
        # Held for review — not applied.
        assert DeviceRepository.find_by_pm(db, "PM-900").location == "Bench"


def test_late_landing_write_keeps_baseline(e2e, tmp_path, monkeypatch):
    """A timed-out write worker can still land later; the tick adopts the
    intended rows as last_write_rows so the late write is not misread as a
    hand edit on the next detection pass."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V", "SN-9", "Bench"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["seeded"] is True
    add_device(h, name="New Tool", pm_number="PM-777", locker_slot=None)

    allow_write = threading.Event()
    real_write = WorkbookAdapter.write_sheet

    def slow_write(self, *args, **kwargs):
        assert allow_write.wait(timeout=10.0)
        return real_write(self, *args, **kwargs)

    monkeypatch.setattr(WorkbookAdapter, "write_sheet", slow_write)
    monkeypatch.setattr(mirror, "_IO_TIMEOUT_SECONDS", 0.3)

    mirror.mark_dirty()
    result = _tick_now()
    assert result["error"] == "timeout"
    # The intended baseline is recorded even though nothing landed yet —
    # it includes PM-777, which the previous baseline does not.
    assert {r[0] for r in mirror._load_state()["last_write_rows"]} == {
        "PM-777", "PM-900",
    }
    assert mirror._load_state()["pending_writes"] is True

    # The abandoned worker lands the write late; the next tick must not flag
    # it as a hand edit.
    allow_write.set()
    monkeypatch.setattr(mirror, "_IO_TIMEOUT_SECONDS", 30.0)
    deadline = time.monotonic() + 5.0
    while mirror._writer_lock.locked() and time.monotonic() < deadline:
        time.sleep(0.05)

    result = _tick_now()
    assert result["error"] is None
    state = mirror._load_state()
    assert state["external_pending"] is False
    assert _mirror_cell(path, "PM-900") == "Bench"


def test_busy_reader_fails_fast(e2e, tmp_path, monkeypatch):
    """While a wedged reader holds the gate, another read reports a timeout
    immediately instead of leaking another worker thread."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-9", "Scope", "Tool", "M", "Mod", "S9", "Bench"],
    ])
    _point_mirror_at(monkeypatch, path)
    e2e()

    assert mirror._reader_lock.acquire(blocking=False)
    try:
        started = time.monotonic()
        parsed, err = mirror._read_catalog(path)
        assert time.monotonic() - started < mirror._IO_TIMEOUT_SECONDS
        assert parsed is None
        assert err == "timeout"
    finally:
        mirror._reader_lock.release()


def test_write_targets_catalog_sheet_not_active(e2e, tmp_path, monkeypatch):
    """A workbook left open on a scratch sheet: the mirror rewrites the
    sheet carrying the PM column, not the active one."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V", "SN-9", "Bench"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["seeded"] is True

    wb = load_workbook(path)
    notes = wb.create_sheet("Notes")
    notes.append(["hand notes", "do not touch"])
    wb.active = wb.sheetnames.index("Notes")
    wb.save(path)
    wb.close()

    mirror.mark_dirty()
    _tick_now()

    # The Notes sheet survives untouched; the catalog sheet was rewritten
    # canonically even though it was not the active sheet.
    wb = load_workbook(path)
    try:
        assert [c.value for c in wb["Notes"][1]] == ["hand notes", "do not touch"]
        rows = list(wb["Sheet"].iter_rows(values_only=True))
        headers = [str(c).strip() if c else "" for c in rows[0]]
        assert headers == MIRROR_HEADERS
        pm_col = headers.index("PM Number")
        assert any(r[pm_col] == "PM-900" for r in rows[1:])
    finally:
        wb.close()


def test_external_diffs_missing_file_reports_via_read(tmp_path, monkeypatch):
    """No raw pre-stat: a missing file surfaces the timeboxed read's own
    'File not found' error through the diffs API."""
    _point_mirror_at(monkeypatch, tmp_path / "mirror.xlsx")
    diffs, err = mirror.external_diffs()
    assert diffs == []
    assert "file not found" in err.lower()


def _edit_location(path, pm_number: str, value: str) -> None:
    """Hand-edit one row's Location cell in the mirror workbook."""
    wb = load_workbook(path)
    try:
        ws = wb.active
        headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        pm_col = headers.index("PM Number") + 1
        loc_col = headers.index("Location") + 1
        for row in ws.iter_rows(min_row=2):
            if str(row[pm_col - 1].value).strip() == pm_number:
                ws.cell(row=row[0].row, column=loc_col, value=value)
                return
        raise AssertionError(f"PM {pm_number} not in sheet {path}")
    finally:
        wb.save(path)
        wb.close()


def test_maintenance_word_on_cabinet_unit_marks_maintenance(
    e2e, tmp_path, monkeypatch
):
    """A hand edit writing the maintenance word over a cabinet unit's
    derived Location does the same thing as the dashboard action."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    device_id = add_device(h, name="Scope", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None
    assert _mirror_cell(path, "PM-100") == in_locker_token()

    # Operator hand-edits the unit's cell to the maintenance word.
    _edit_location(path, "PM-100", "maintenance")
    result = _tick_now()
    assert result["external"] is True

    outcome = mirror.apply_external(get_engine())
    assert outcome["applied"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.MAINTENANCE

    # The next flush writes the canonical token — sheet and DB converge.
    _drain_mirror()
    assert _mirror_cell(path, "PM-100") == maintenance_token()


def test_maintenance_word_skipped_while_borrowed(
    e2e, tmp_path, monkeypatch
):
    """The maintenance word is refused while the unit is borrowed — the
    loan owns the row, and the mirror rewrites the borrower on flush."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    device_id = add_device(
        h, name="Scope", pm_number="PM-100", locker_slot=1, status="borrowed"
    )

    result = _tick_now()
    assert result["error"] is None

    _edit_location(path, "PM-100", "Maintenance")
    result = _tick_now()
    assert result["external"] is True

    outcome = mirror.apply_external(get_engine())
    assert outcome["skipped"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.BORROWED


def test_adoption_honors_maintenance_word_on_registered_row(
    e2e, tmp_path, monkeypatch
):
    """First-sight adoption: a registered unit whose sheet Location says
    the maintenance word adopts into maintenance, not availability."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Scope", "Tool", "M", "Mod", "SN-1", "maintenance"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    device_id = add_device(h, name="Scope", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None
    assert result["seeded"] is True
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.MAINTENANCE
    # The regenerated sheet writes the canonical token.
    assert _mirror_cell(path, "PM-100") == maintenance_token()


def test_apply_external_missing_file_raises_unavailable(
    e2e, tmp_path, monkeypatch
):
    _no_scheduler(monkeypatch)
    e2e()
    _point_mirror_at(monkeypatch, tmp_path / "mirror.xlsx")
    with pytest.raises(mirror.MirrorUnavailable, match="File not found"):
        mirror.apply_external(get_engine())


def test_apply_external_changed_commit_integrity_error_skips(
    e2e, tmp_path, monkeypatch
):
    """A commit that raises IntegrityError on a 'changed' diff is skipped —
    the remaining diffs still apply."""
    _no_scheduler(monkeypatch)
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-1", "Alpha", "Tool", "M", "Mod", "S1", "Bench"],
        ["PM-2", "Beta", "Tool", "M", "Mod", "S2", "Shelf"],
    ])
    _point_mirror_at(monkeypatch, path)
    h = e2e()

    result = _tick_now()
    assert result["seeded"] is True

    wb = load_workbook(path)
    ws = wb.active
    name_col = MIRROR_HEADERS.index("Name") + 1
    ws.cell(row=2, column=name_col, value="Alpha2")
    ws.cell(row=3, column=name_col, value="Beta2")
    wb.save(path)
    wb.close()
    result = _tick_now()
    assert result["external"] is True

    real_commit = Session.commit
    calls = {"n": 0}

    def flaky_commit(session):
        calls["n"] += 1
        if calls["n"] == 1:
            raise IntegrityError("UPDATE devices", {}, Exception("unique"))
        return real_commit(session)

    monkeypatch.setattr(Session, "commit", flaky_commit)
    try:
        outcome = mirror.apply_external(get_engine())
    finally:
        mirror.flush_scheduled()
    # PM-1's diff sorts first and hits the failing commit; PM-2 applies.
    assert outcome == {"applied": 1, "skipped": 1}
    with h.db() as db:
        assert DeviceRepository.find_by_pm(db, "PM-1").name == "Alpha"
        assert DeviceRepository.find_by_pm(db, "PM-2").name == "Beta2"
    assert mirror._load_state()["external_pending"] is False


def test_last_write_snapshot_categorizes_error():
    """/api/health exposure: the write snapshot must store the error
    category, never the raw adapter text (which carries paths)."""
    mirror._remember_write({
        "flushed": False,
        "written": 0,
        "error": "File not found: /mnt/locker/share/smart_locker_catalog.xlsx",
    })
    snap = mirror.last_write()
    assert snap["error"] == "missing"
    assert "/mnt" not in (snap["error"] or "")


def test_added_diff_maintenance_word_on_registered_row(
    e2e, tmp_path, monkeypatch
):
    """A unit registered after the last baseline write is absent from
    last_write_rows, so its hand-written sheet row diffs as "added" — and
    the maintenance word in its Location must still transition the unit."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None
    assert _mirror_cell(path, "PM-100") == in_locker_token()

    # Registered in SQLite after the baseline write: not in
    # last_write_rows, so a sheet row for it is an "added" diff. The sheet
    # rewrite must carry every baseline row plus the new one.
    device_id = add_device(h, name="Scope", pm_number="PM-500", locker_slot=2)
    _write_sheet(path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "", "", "", in_locker_token()],
        ["PM-500", "Scope", "Tool", "", "", "", maintenance_token()],
    ])
    result = _tick_now()
    assert result["external"] is True

    outcome = mirror.apply_external(get_engine())
    assert outcome["applied"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.MAINTENANCE

    # The converged sheet writes the canonical token on the new row.
    _drain_mirror()
    assert _mirror_cell(path, "PM-500") == maintenance_token()


def test_added_diff_maintenance_word_skipped_while_borrowed(
    e2e, tmp_path, monkeypatch
):
    """The "added" branch honors the same borrow guard as "changed": a
    unit registered and loaned after the baseline write stays borrowed —
    the loan owns the row, the sheet word cannot park it."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    user_id = add_user(h, "0A30000001", display_name="E2E Borrower")
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None

    device_id = add_device(
        h, name="Scope", pm_number="PM-500", locker_slot=2, status="borrowed"
    )
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        device.current_borrower_id = user_id
        db.commit()

    _write_sheet(path, [
        CATALOG_HEADERS,
        ["PM-100", "Camera", "Tool", "", "", "", in_locker_token()],
        ["PM-500", "Scope", "Tool", "", "", "", maintenance_token()],
    ])
    result = _tick_now()
    assert result["external"] is True

    outcome = mirror.apply_external(get_engine())
    assert outcome["skipped"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.BORROWED
        assert device.current_borrower_id == user_id


def test_apply_external_flush_error_is_skipped_not_raised(
    e2e, tmp_path, monkeypatch
):
    """A flush-time failure inside _apply_field (OperationalError from a
    locked SQLite/tired SD, CatalogError from a raced borrow) is a skipped
    diff — apply_external must not raise, or the route turns it into a
    bare 500 and external_pending stays set."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    device_id = add_device(h, name="Scope", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None
    assert _mirror_cell(path, "PM-100") == in_locker_token()

    _edit_location(path, "PM-100", "maintenance")
    result = _tick_now()
    assert result["external"] is True

    def broken_place_word(session, device, text):
        raise OperationalError("x", None, None)

    monkeypatch.setattr(mirror, "apply_place_word", broken_place_word)
    try:
        outcome = mirror.apply_external(get_engine())
    finally:
        mirror.flush_scheduled()
    assert outcome["skipped"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.AVAILABLE


@pytest.mark.parametrize("word", ["MAINTENANCE", " maintenance "])
def test_maintenance_word_case_and_padding_apply(
    e2e, tmp_path, monkeypatch, word
):
    """place_kind strips and case-folds the cell: an uppercase or padded
    spelling of the maintenance word still applies the transition."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    device_id = add_device(h, name="Scope", pm_number="PM-100", locker_slot=1)

    result = _tick_now()
    assert result["error"] is None
    assert _mirror_cell(path, "PM-100") == in_locker_token()

    _edit_location(path, "PM-100", word)
    result = _tick_now()
    assert result["external"] is True

    outcome = mirror.apply_external(get_engine())
    assert outcome["applied"] >= 1
    with h.db() as db:
        device = DeviceRepository.find_by_id(db, device_id)
        assert device.status == DeviceStatus.MAINTENANCE


def test_apply_external_via_dashboard_route(e2e, tmp_path, monkeypatch):
    """The HTTP wiring: POST /api/dashboard/mirror/apply runs apply_external
    behind the admin-secret gate — 401 without the header, the same apply
    counts with it."""
    _no_scheduler(monkeypatch)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)
    monkeypatch.setenv("SMART_LOCKER_DASHBOARD_ADMIN_SECRET", "e2e-secret")

    result = _tick_now()
    assert result["error"] is None

    _edit_location(path, "PM-100", "maintenance")
    result = _tick_now()
    assert result["external"] is True

    # The gate is the header, not the button — no header, no apply.
    resp = h.client.post("/api/dashboard/mirror/apply")
    assert resp.status_code == 401

    resp = h.client.post(
        "/api/dashboard/mirror/apply",
        headers=dashboard_admin_headers("e2e-secret"),
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["ok"] is True
    assert body["applied"] >= 1
    mirror.flush_scheduled()
    with h.db() as db:
        device = DeviceRepository.find_by_pm(db, "PM-100")
        assert device.status == DeviceStatus.MAINTENANCE
