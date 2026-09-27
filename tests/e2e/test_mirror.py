"""
File: test_mirror.py
Description: End-to-end coverage of the catalog mirror — the one sync path
             the API suite cannot fully exercise. Boots create_app() via the
             ``e2e`` factory against a real workbook on disk: first sight of
             a populated sheet adopts it into the SQLite catalog, taps drive
             borrow/return through the NFC bridge and mark the mirror dirty,
             and flush_scheduled() drains the async writer before the xlsx
             is reopened and the Location cell asserted.
Project: smart_locker/tests/e2e
Notes: SMART_LOCKER_MIRROR_PATH is pointed at a temp workbook per test — the
       suite autouse fixture leaves the mirror unconfigured. Workbooks built
       with tests.api.helpers.catalog_workbook; the mirror regenerates the
       sheet with MIRROR_HEADERS, so header order after a flush is canonical.
"""

import threading
import time

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from config.settings import in_locker_token, maintenance_token
from smart_locker.database.engine import get_engine
from smart_locker.database.models import DeviceStatus
from smart_locker.database.repositories import DeviceRepository, RegistrantRepository
from smart_locker.sync import mirror, scheduler, sync_status
from smart_locker.sync.catalog_sheet import MIRROR_HEADERS
from tests.api.helpers import catalog_workbook, dashboard_admin_headers
from tests.e2e.helpers import add_device, add_user, get_device, uid_hmac_for

# Simulated UIDs — hex-like strings. The DB is fresh per test, so constants
# can repeat safely across tests.
CARD_USER = "0A30000001"
CARD_ADMIN = "0A30000002"
TAG_PM100 = "0A40000001"
TAG_PM999 = "0A40000002"

CATALOG_HEADERS = ["PM", "Name", "Type", "Manufacturer", "Model", "Serial", "Location"]


def _login(h, card_uid: str) -> dict:
    """Tap a work card and wait for the session to open."""
    h.tap(card_uid)
    return h.wait_event("auth_success")


def _admin_session(h) -> dict:
    """Enroll an admin and open the loopback admin-panel session."""
    add_user(h, CARD_ADMIN, display_name="E2E Admin", role="admin")
    resp = h.client.post("/api/admin/session")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["success"] is True
    return body


def _point_mirror_at(monkeypatch, path) -> None:
    """Aim the mirror at this test's workbook (overrides the autouse blank)."""
    monkeypatch.setenv("SMART_LOCKER_MIRROR_PATH", str(path))


def _mirror_cell(path, pm_number: str, header: str = "Location"):
    """Reopen the written mirror and return one cell for the PM row."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(c).strip() if c else "" for c in rows[0]]
    pm_col, want_col = headers.index("PM Number"), headers.index(header)
    for row in rows[1:]:
        if row[pm_col] is not None and str(row[pm_col]).strip() == pm_number:
            return None if row[want_col] is None else str(row[want_col]).strip()
    raise AssertionError(f"PM {pm_number} not in mirror {path}")


def _wait_for_trigger(trigger: str, timeout: float) -> dict:
    """Poll sync_status.get() until the last tick ran under ``trigger``."""
    deadline = time.monotonic() + timeout
    last: dict = {}
    while time.monotonic() < deadline:
        snap = sync_status.get()
        if snap.get("at") and snap.get("trigger") == trigger:
            return snap
        last = snap
        time.sleep(0.05)
    raise TimeoutError(f"No '{trigger}' tick within {timeout}s. Last: {last}")


def _tick_now(timeout: float = 15.0):
    """Run one mirror tick, waiting out the app's startup tick if needed.

    The lifespan scheduler queues a startup tick on a worker thread — a
    manual or scheduled tick that lands while it runs is skipped as
    ``in_progress``. Production defers to the next interval tick; tests
    retry so the write is deterministic.
    """
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        result = mirror.tick(get_engine(), trigger="manual")
        if result.get("skipped") != "in_progress":
            return result
        time.sleep(0.1)
    raise TimeoutError("mirror tick never became free")


def _drain_mirror(timeout: float = 15.0) -> None:
    """Join scheduled flush workers, then run the tick they may have skipped.

    ``schedule_flush`` workers that collide with an in-flight tick defer to
    the next interval — tests run the same follow-up tick explicitly so the
    pending write lands now instead of up to a minute later.
    """
    mirror.flush_scheduled(timeout=timeout)
    _tick_now(timeout=timeout)


def _sync_source(h, timeout: float = 15.0) -> dict:
    """POST the manual sync, retrying while the startup tick holds the mutex."""
    deadline = time.monotonic() + timeout
    while True:
        resp = h.client.post("/api/admin/sync-source")
        if resp.status_code != 409:
            return resp
        if time.monotonic() >= deadline:
            raise TimeoutError("sync-source stayed 409 — tick never released")
        time.sleep(0.1)


def test_first_tick_adopts_sheet_then_writes_mirror(e2e, tmp_path, monkeypatch):
    """First sight of a populated sheet adopts it: unknown PMs insert as
    catalog rows, an existing locker row keeps slot/tag while its catalog
    fields refresh, Location person names seed registrants — then the same
    tick regenerates the sheet canonically with derived Locations."""
    h = e2e()
    device_id = add_device(
        h, name="Old Scope", pm_number="PM-100", locker_slot=5,
        tag_uid=TAG_PM100,
    )
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-100", "Rigol DS1054", "Oscilloscope", "Rigol", "DS1054Z",
         "SN-RIGOL-1", "Locker"],
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Alice Smith"],
    ])
    _point_mirror_at(monkeypatch, path)
    _admin_session(h)

    resp = _sync_source(h)
    assert resp.status_code == 200, resp.text
    assert resp.json()["success"] is True
    # Whether the startup tick or this POST adopted first, the catalog now
    # holds the sheet's rows — timing decides which one ran the adopt.
    _drain_mirror()

    # Existing locker row: catalog fields refreshed, locker fields kept.
    device = get_device(h, device_id)
    assert device.name == "Rigol DS1054"
    assert device.serial_number == "SN-RIGOL-1"
    assert device.status == DeviceStatus.AVAILABLE
    assert device.locker_slot == 5
    assert device.tag_hmac == uid_hmac_for(TAG_PM100)

    # New PMs are now SQLite catalog rows — the one-time catalog migration.
    with h.db() as db:
        adopted = DeviceRepository.find_by_pm(db, "PM-900")
        assert adopted is not None
        assert adopted.locker_slot is None
        assert adopted.location == "Alice Smith"
        registrants = {r.display_name for r in RegistrantRepository.get_all(db)}
    assert "Alice Smith" in registrants

    # The mirror was regenerated canonically: headers and derived Location.
    assert _mirror_cell(path, "PM-100") == in_locker_token()
    assert _mirror_cell(path, "PM-900") == "Alice Smith"

    # Recorded for the admin footer and health probe.
    status = h.client.get("/api/admin/sync-status").json()
    assert status["trigger"] == "manual"
    assert status["mirror"]["configured"] is True


def test_borrow_writes_borrower_name_and_creates_file(e2e, tmp_path, monkeypatch):
    """A sticker borrow marks the mirror dirty; the flush writes the whole
    catalog — creating the file when it does not exist yet."""
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    user_id = add_user(h, CARD_USER, display_name="E2E Borrower")
    device_id = add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)
    payload = h.wait_event("device_action")
    assert payload["success"] is True
    assert payload["action"] == "borrow"
    assert get_device(h, device_id).current_borrower_id == user_id

    _drain_mirror()
    assert path.exists()
    assert _mirror_cell(path, "PM-100") == "E2E Borrower"


def test_failed_borrow_commit_does_not_mirror_uncommitted_loan(
    e2e, tmp_path, monkeypatch
):
    """A failed HTTP commit leaves the sheet at its last written state."""
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    device_id = add_device(h, pm_number="PM-100", tag_uid=TAG_PM100)
    _login(h, CARD_USER)

    # Baseline: seed the mirror so it shows the in-locker token first.
    mirror.mark_dirty()
    _tick_now()
    assert _mirror_cell(path, "PM-100") == in_locker_token()

    real_commit = Session.commit

    def fail_loan_commit(session):
        if session.info.get("mirror_dirty_pending"):
            raise RuntimeError("simulated SQLite commit failure")
        return real_commit(session)

    monkeypatch.setattr(Session, "commit", fail_loan_commit)
    with pytest.raises(RuntimeError, match="simulated SQLite commit failure"):
        h.client.post(f"/api/devices/{device_id}/borrow")

    _drain_mirror()
    assert get_device(h, device_id).current_borrower_id is None
    assert _mirror_cell(path, "PM-100") == in_locker_token()


def test_return_writes_in_locker_token(e2e, tmp_path, monkeypatch):
    """A return puts the in-locker token (not a blank) back into Location."""
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)
    borrow = h.wait_event("device_action")
    assert borrow["action"] == "borrow"
    _drain_mirror()
    assert _mirror_cell(path, "PM-100") == "E2E Borrower"

    h.tap(TAG_PM100)
    returned = h.wait_event("device_action")
    assert returned["success"] is True
    assert returned["action"] == "return"

    _drain_mirror()
    assert _mirror_cell(path, "PM-100") == in_locker_token()


def test_maintenance_device_gets_maintenance_token(e2e, tmp_path, monkeypatch):
    """The Pi rewrites the whole sheet: a maintenance unit's Location is the
    maintenance token, not whatever text the sheet used to carry."""
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="In Repair", pm_number="PM-999", locker_slot=2,
        tag_uid=TAG_PM999, status="maintenance",
    )

    # Any catalog change flushes the whole sheet — maintenance included.
    mirror.mark_dirty()
    _tick_now()
    assert _mirror_cell(path, "PM-999") == maintenance_token()


def test_health_reports_last_writeback(e2e, tmp_path, monkeypatch):
    """GET /api/health exposes the mirror-write snapshot after a flush."""
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)
    h = e2e()
    add_user(h, CARD_USER, display_name="E2E Borrower")
    add_device(
        h, name="Camera", pm_number="PM-100", locker_slot=1, tag_uid=TAG_PM100
    )

    _login(h, CARD_USER)
    h.tap(TAG_PM100)
    h.wait_event("device_action")
    _drain_mirror()

    resp = h.client.get("/api/health")
    assert resp.status_code == 200
    wb = resp.json()["last_writeback"]
    assert wb is not None
    assert wb["error"] is None
    assert wb["saved"] is True
    assert wb["written"] >= 1
    assert wb["at"]


def test_register_promotes_adopted_catalog_row(e2e, tmp_path, monkeypatch):
    """An adopted sheet row is registered into a slot; the mirror then shows
    the in-locker token on that row."""
    h = e2e()
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-200", "Multimeter", "Tool", "Fluke", "87V", "SN-200", "Locker"],
    ])
    _point_mirror_at(monkeypatch, path)
    _admin_session(h)

    # Adopt the sheet so PM-200 exists in the catalog, then register it.
    resp = _sync_source(h)
    assert resp.status_code == 200
    resp = h.client.post(
        "/api/admin/devices/register",
        json={"pm_number": "PM-200", "locker_slot": 7},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["success"] is True
    assert body["locker_slot"] == 7

    assert h.ctx.pending_tag_bind is not None
    device = get_device(h, h.ctx.pending_tag_bind.device_id)
    assert device.pm_number == "PM-200"
    assert device.locker_slot == 7
    assert device.status is DeviceStatus.AVAILABLE

    _drain_mirror()
    assert _mirror_cell(path, "PM-200") == in_locker_token()


def test_sync_source_409_when_tick_running(e2e, tmp_path, monkeypatch):
    """Holding the real tick mutex maps to HTTP 409 for the manual sync."""
    h = e2e()
    _admin_session(h)

    assert mirror._tick_lock.acquire(blocking=False)
    try:
        resp = h.client.post("/api/admin/sync-source")
        assert resp.status_code == 409
        assert "already running" in resp.json()["detail"].lower()
    finally:
        mirror._tick_lock.release()


def test_sync_source_unconfigured_mirror_is_noop(e2e):
    """No mirror configured is a skipped tick, not a kiosk error."""
    h = e2e()
    _admin_session(h)
    resp = h.client.post("/api/admin/sync-source")
    assert resp.status_code == 200
    assert resp.json()["skipped"] == "unconfigured"


def test_hand_edit_gated_then_applied(e2e, tmp_path, monkeypatch):
    """A hand edit in the file is held for review; the admin Apply lands it
    in SQLite. The sheet is never silently merged."""
    h = e2e()
    path = catalog_workbook(tmp_path, [
        CATALOG_HEADERS,
        ["PM-900", "Fluke 87V", "Multimeter", "Fluke", "87V",
         "SN-FLUKE-9", "Workshop"],
    ])
    _point_mirror_at(monkeypatch, path)
    _admin_session(h)
    monkeypatch.setenv("SMART_LOCKER_DASHBOARD_ADMIN_SECRET", "e2e-secret")

    # Adopt + flush baseline.
    _tick_now()
    assert _mirror_cell(path, "PM-900") == "Workshop"

    # Hand-edit Location, then tick: the change is flagged, not applied.
    wb = load_workbook(path)
    ws = wb.active
    loc_col = MIRROR_HEADERS.index("Location") + 1
    ws.cell(row=2, column=loc_col, value="Moved by hand")
    wb.save(path)
    wb.close()

    result = _tick_now()
    assert result["external"] is True
    with h.db() as db:
        assert DeviceRepository.find_by_pm(db, "PM-900").location == "Workshop"

    diffs = h.client.get(
        "/api/dashboard/mirror/diffs",
        headers=dashboard_admin_headers("e2e-secret"),
    ).json()["diffs"]
    assert any(d["kind"] == "changed" for d in diffs)

    resp = h.client.post(
        "/api/dashboard/mirror/apply",
        headers=dashboard_admin_headers("e2e-secret"),
    )
    assert resp.status_code == 200
    assert resp.json()["applied"] >= 1
    with h.db() as db:
        assert DeviceRepository.find_by_pm(db, "PM-900").location == "Moved by hand"

    _drain_mirror()
    assert h.client.get("/api/dashboard/mirror").json()["external_changes"] is False


def test_start_scheduler_runs_immediate_startup_tick(e2e, tmp_path, monkeypatch):
    """start_scheduler queues one startup tick that writes the mirror."""
    h = e2e()
    add_device(h, name="Camera", pm_number="PM-100", locker_slot=1)
    path = tmp_path / "mirror.xlsx"
    _point_mirror_at(monkeypatch, path)

    try:
        scheduler.start_scheduler(get_engine(), interval_seconds=3600)
        snap = _wait_for_trigger("startup", timeout=5.0)
        assert snap["ok"] is True
        deadline = time.monotonic() + 5.0
        while not path.exists() and time.monotonic() < deadline:
            time.sleep(0.05)
        assert _mirror_cell(path, "PM-100") == in_locker_token()
    finally:
        scheduler.stop_scheduler()


def test_health_available_while_startup_tick_blocked(e2e, tmp_path, monkeypatch):
    """The real lifespan exposes health while its startup tick is blocked."""
    tick_started = threading.Event()
    allow_tick_to_finish = threading.Event()

    def slow_tick(engine, trigger="interval"):
        assert trigger == "startup"
        tick_started.set()
        assert allow_tick_to_finish.wait(timeout=5.0)
        return {}

    monkeypatch.setattr(scheduler, "run_mirror_tick", slow_tick)
    monkeypatch.setattr(scheduler, "_try_dashboard_launcher", lambda: None)

    try:
        h = e2e()
        assert h.client.get("/api/health").status_code == 200
        assert tick_started.wait(timeout=1.0)
    finally:
        allow_tick_to_finish.set()
        scheduler.stop_scheduler()
