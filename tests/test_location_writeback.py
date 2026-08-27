"""
File: test_location_writeback.py
Description: Tests for Pi → Excel write-back of Location only.
             Locker available → in-locker token; borrowed → borrower name.
             A locked or missing workbook must not raise into the kiosk.
Project: smart_locker/tests
Notes: Run with: python -m pytest tests/test_location_writeback.py -v
       Uses temporary workbooks; no NFC hardware.
"""

import os
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from smart_locker.database.models import DeviceStatus
from smart_locker.database.repositories import DeviceRepository, UserRepository
from smart_locker.security.encryption import encrypt
from smart_locker.security.hashing import compute_uid_hmac
from smart_locker.services.device_registration import register_locker_device
from smart_locker.services.locker_service import LockerService
from smart_locker.sync.location_writeback import (
    IN_LOCKER_TOKEN,
    last_writeback,
    write_location,
    write_location_with_engine,
)
from smart_locker.sync.scheduler import _run_source_import


def _workbook(path: Path, rows: list[list], extra_sheet: str | None = None) -> Path:
    """Write a workbook (first row = headers) and return ``path``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    for row in rows:
        ws.append(row)
    if extra_sheet:
        other = wb.create_sheet(extra_sheet)
        other.append(["Keep", "Me"])
        other.append(["untouched", 1])
    wb.save(path)
    return path


def _location_by_pm(path: Path) -> dict[str, str]:
    """Read Equipment → Location from the first sheet."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    headers = [str(h).strip() if h else "" for h in rows[0]]
    pm_i = headers.index("Equipment")
    loc_i = headers.index("Location")
    out: dict[str, str] = {}
    for row in rows[1:]:
        pm = str(row[pm_i]).strip() if row[pm_i] is not None else ""
        loc = "" if row[loc_i] is None else str(row[loc_i]).strip()
        if pm:
            out[pm] = loc
    return out


def _add_user(db_session, enc_key, hmac_key, name: str = "Alice"):
    """Insert one enrolled user and return it."""
    uid = "A1B2C3D4"
    return UserRepository.create(
        db_session,
        display_name=name,
        uid_hmac=compute_uid_hmac(uid, hmac_key),
        encrypted_card_uid=encrypt(uid, enc_key),
    )


class TestWriteLocation:
    """write_location maps locker SQLite state onto Location by PM."""

    def test_available_writes_in_locker_token(self, db_session, tmp_path):
        """An available locker device is written as the stable in-locker token."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer", "Location"],
            ["PM-001", "Fluke", "Old Name"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            manufacturer="Fluke",
            locker_slot=1,
            status=DeviceStatus.AVAILABLE.value,
        )
        db_session.flush()

        result = write_location(db_session, path)

        assert result.written == 1
        assert result.saved is True
        assert _location_by_pm(path)["PM-001"] == IN_LOCKER_TOKEN

    def test_borrowed_writes_borrower_name(
        self, db_session, tmp_path, enc_key, hmac_key
    ):
        """A borrowed locker device is written as the borrower's display name."""
        user = _add_user(db_session, enc_key, hmac_key, "Alice")
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer", "Location"],
            ["PM-002", "Keysight", "Locker"],
        ])
        DeviceRepository.create(
            db_session,
            name="Scope",
            device_type="general",
            pm_number="PM-002",
            manufacturer="Keysight",
            locker_slot=2,
            status=DeviceStatus.BORROWED.value,
            current_borrower_id=user.id,
        )
        db_session.flush()

        result = write_location(db_session, path)

        assert result.written == 1
        assert _location_by_pm(path)["PM-002"] == "Alice"

    def test_formula_like_location_stored_as_text(self, tmp_path):
        """Location starting with =+@- is stored as text, not an Excel formula."""
        from smart_locker.sync.location_writeback import write_location_value

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Name", "Location"],
            ["PM-VAN", "Van kit", "Workshop"],
        ])
        evil = '=HYPERLINK("http://evil.example/x","x")'
        result = write_location_value(path, "PM-VAN", evil)
        assert result.error is None
        assert result.saved is True
        wb = load_workbook(path, data_only=False)
        try:
            cell = wb.active["C2"]
            assert cell.data_type == "s"
            stored = "" if cell.value is None else str(cell.value)
            assert stored.lstrip("'") == evil
            assert cell.data_type != "f"
        finally:
            wb.close()

    def test_preserves_other_columns_and_sheets(self, db_session, tmp_path):
        """Only Location changes; catalog cells and extra sheets stay."""
        path = _workbook(
            tmp_path / "device-list.xlsx",
            [
                ["Equipment", "Manufacturer", "Model", "Location"],
                ["PM-001", "Fluke", "87V", "Someone"],
            ],
            extra_sheet="Notes",
        )
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            manufacturer="Fluke",
            model="87V",
            locker_slot=1,
        )
        db_session.flush()

        write_location(db_session, path)

        wb = load_workbook(path)
        try:
            ws = wb.active
            assert ws["B2"].value == "Fluke"
            assert ws["C2"].value == "87V"
            assert ws["D2"].value == IN_LOCKER_TOKEN
            assert "Notes" in wb.sheetnames
            notes = wb["Notes"]
            assert notes["A2"].value == "untouched"
            assert notes["B2"].value == 1
        finally:
            wb.close()

    def test_non_locker_pm_is_left_alone(self, db_session, tmp_path):
        """Excel rows that are not locker devices keep their Location."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-LOCKER", "Stale"],
            ["PM-FIELD", "Bob"],
        ])
        DeviceRepository.create(
            db_session,
            name="Locker unit",
            device_type="general",
            pm_number="PM-LOCKER",
            locker_slot=1,
        )
        db_session.flush()

        write_location(db_session, path)

        cells = _location_by_pm(path)
        assert cells["PM-LOCKER"] == IN_LOCKER_TOKEN
        assert cells["PM-FIELD"] == "Bob"

    def test_locker_pm_missing_from_excel_does_not_add_a_row(
        self, db_session, tmp_path
    ):
        """Write-back never inserts Excel rows for locker PMs the sheet lacks."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-OTHER", "Desk"],
        ])
        DeviceRepository.create(
            db_session,
            name="Ghost",
            device_type="general",
            pm_number="PM-MISSING",
            locker_slot=3,
        )
        db_session.flush()

        result = write_location(db_session, path)

        assert result.skipped >= 1
        wb = load_workbook(path, read_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()
        assert len(rows) == 2
        assert _location_by_pm(path)["PM-OTHER"] == "Desk"

    def test_unchanged_location_does_not_rewrite(self, db_session, tmp_path):
        """When every locker cell already matches, the file mtime is left alone."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", IN_LOCKER_TOKEN],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        os.utime(path, (1_700_000_000, 1_700_000_000))
        before = path.stat().st_mtime

        result = write_location(db_session, path)

        assert result.written == 0
        assert result.saved is False
        assert path.stat().st_mtime == before


class TestWritebackResilience:
    """A missing, locked, or incomplete workbook must not raise."""

    def test_missing_file_does_not_raise(self, db_session, tmp_path):
        """Share down / wrong path is skipped, not fatal."""
        result = write_location(db_session, tmp_path / "no-such.xlsx")
        assert result.saved is False
        snap = last_writeback()
        assert snap is not None
        assert snap["error"] == "missing"
        assert snap["saved"] is False

    def test_missing_location_column_does_not_raise(self, db_session, tmp_path):
        """A sheet without Location is skipped, not rewritten."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer"],
            ["PM-001", "Fluke"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        mtime = path.stat().st_mtime

        result = write_location(db_session, path)

        assert result.saved is False
        assert path.stat().st_mtime == mtime

    def test_retries_when_catalog_changes_during_write(self, db_session, tmp_path):
        """A catalog save during copy/edit is not reverted; Location still updates."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer", "Location"],
            ["PM-001", "Fluke", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            manufacturer="Fluke",
            locker_slot=1,
        )
        db_session.flush()

        import shutil as shutil_mod

        original_copy = shutil_mod.copy2
        calls = {"n": 0}

        def copy_then_edit_catalog(src, dst, *args, **kwargs):
            original_copy(src, dst, *args, **kwargs)
            calls["n"] += 1
            if calls["n"] == 1:
                wb = load_workbook(src)
                try:
                    wb.active["B2"] = "Keysight"
                    wb.save(src)
                finally:
                    wb.close()

        with patch(
            "smart_locker.sync.location_writeback.shutil.copy2",
            copy_then_edit_catalog,
        ):
            result = write_location(db_session, path)

        assert result.saved is True
        wb = load_workbook(path)
        try:
            assert wb.active["B2"].value == "Keysight"
            assert wb.active["C2"].value == IN_LOCKER_TOKEN
        finally:
            wb.close()

    def test_locked_replace_does_not_raise(self, db_session, tmp_path):
        """Excel holding the file open: log and skip, never raise."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()

        with patch("smart_locker.sync.location_writeback.time.sleep"):
            with patch.object(Path, "replace", side_effect=PermissionError("locked")):
                result = write_location(db_session, path)

        assert result.saved is False
        assert _location_by_pm(path)["PM-001"] == "Old"

    def test_retries_then_succeeds(self, db_session, tmp_path):
        """A momentary lock is retried; the cell is written on a later attempt."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        original = Path.replace
        calls = {"n": 0}

        def flaky(self, target):
            calls["n"] += 1
            if calls["n"] == 1:
                raise PermissionError("locked")
            return original(self, target)

        with patch("smart_locker.sync.location_writeback.time.sleep"):
            with patch.object(Path, "replace", flaky):
                result = write_location(db_session, path)

        assert result.saved is True
        assert _location_by_pm(path)["PM-001"] == IN_LOCKER_TOKEN


class TestWritebackWiring:
    """Borrow/return, Register Device, and scheduled sync all write Location."""

    def test_borrow_writes_borrower_name(
        self, db_session, tmp_path, enc_key, hmac_key, monkeypatch
    ):
        """A successful kiosk borrow updates Location to the user."""
        from smart_locker.auth.session_manager import SessionManager

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", IN_LOCKER_TOKEN],
        ])
        monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))
        user = _add_user(db_session, enc_key, hmac_key, "Alice")
        device = DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        session = SessionManager(timeout_seconds=60).start_session(user)

        assert LockerService.borrow_device(db_session, session, device.id) is True
        from smart_locker.sync.location_writeback import flush_scheduled_writeback

        flush_scheduled_writeback()
        assert _location_by_pm(path)["PM-001"] == "Alice"

    def test_return_writes_in_locker_token(
        self, db_session, tmp_path, enc_key, hmac_key, monkeypatch
    ):
        """Returning a device writes the in-locker token back into Excel."""
        from smart_locker.auth.session_manager import SessionManager

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Alice"],
        ])
        monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))
        user = _add_user(db_session, enc_key, hmac_key, "Alice")
        device = DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
            status=DeviceStatus.BORROWED.value,
            current_borrower_id=user.id,
        )
        db_session.flush()
        session = SessionManager(timeout_seconds=60).start_session(user)

        assert LockerService.return_device(db_session, session, device.id) is True
        from smart_locker.sync.location_writeback import flush_scheduled_writeback

        flush_scheduled_writeback()
        assert _location_by_pm(path)["PM-001"] == IN_LOCKER_TOKEN

    def test_register_device_writes_in_locker_token(self, db_session, tmp_path):
        """A newly registered locker PM is written as in-locker in Excel."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer", "Model", "Location"],
            ["PM-NEW", "Fluke", "87V", "Desk"],
        ])

        register_locker_device(db_session, str(path), "PM-NEW", 4)

        assert _location_by_pm(path)["PM-NEW"] == IN_LOCKER_TOKEN

    def test_scheduler_import_writes_location(self, db_session, tmp_path):
        """Startup/interval import is followed by Location write-back."""
        from smart_locker.database.engine import get_engine

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Manufacturer", "Model", "Location"],
            ["PM-SCHED-001", "NewMfr", "NewModel", "Stale Person"],
        ])
        DeviceRepository.create(
            db_session,
            name="Old",
            device_type="general",
            pm_number="PM-SCHED-001",
            manufacturer="OldMfr",
            model="OldModel",
            locker_slot=1,
        )
        db_session.commit()

        _run_source_import(get_engine(), path)

        device = DeviceRepository.find_by_pm(db_session, "PM-SCHED-001")
        assert device is not None
        assert device.manufacturer == "NewMfr"
        assert _location_by_pm(path)["PM-SCHED-001"] == IN_LOCKER_TOKEN

    def test_engine_helper_matches_session_write(self, db_session, tmp_path):
        """write_location_with_engine sees committed locker state."""
        from smart_locker.database.engine import get_engine

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.commit()

        result = write_location_with_engine(get_engine(), path)
        assert result.saved is True
        assert _location_by_pm(path)["PM-001"] == IN_LOCKER_TOKEN


class TestSiteWritebackAliases:
    """Write-back uses the same env header lists and in-locker token as import."""

    def test_id_header_extra_finds_join_column(self, db_session, tmp_path, monkeypatch):
        """SMART_LOCKER_ID_HEADERS lets write-back match a non-default ID column."""
        monkeypatch.setenv("SMART_LOCKER_ID_HEADERS", "Inventory No")
        path = _workbook(tmp_path / "inventory.xlsx", [
            ["Inventory No", "Location"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        result = write_location(db_session, path)
        assert result.saved is True
        wb = load_workbook(path)
        try:
            assert wb.active["B2"].value == IN_LOCKER_TOKEN
        finally:
            wb.close()

    def test_location_header_extra_finds_location_column(
        self, db_session, tmp_path, monkeypatch
    ):
        """SMART_LOCKER_LOCATION_HEADERS lets write-back find Location under another name."""
        monkeypatch.setenv("SMART_LOCKER_LOCATION_HEADERS", "Whereabouts")
        path = _workbook(tmp_path / "inventory.xlsx", [
            ["Equipment", "Whereabouts"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        result = write_location(db_session, path)
        assert result.saved is True
        wb = load_workbook(path)
        try:
            assert wb.active["B2"].value == IN_LOCKER_TOKEN
        finally:
            wb.close()

    def test_custom_in_locker_token_written(self, db_session, tmp_path, monkeypatch):
        """Available devices write SMART_LOCKER_IN_LOCKER_TOKEN, not the default Locker."""
        monkeypatch.setenv("SMART_LOCKER_IN_LOCKER_TOKEN", "At base")
        path = _workbook(tmp_path / "inventory.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Old"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        result = write_location(db_session, path)
        assert result.saved is True
        assert _location_by_pm(path)["PM-001"] == "At base"


class TestWritebackColumnAndStatus:
    """I2 location-over-owner, I20/I21 join keys, I23 MAINTENANCE skip."""

    def test_owner_and_location_updates_location_only(self, db_session, tmp_path):
        """A sheet with Owner left of Location writes Location, not Owner."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Owner", "Location"],
            ["PM-001", "Alice", "Stale"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        write_location(db_session, path)
        wb = load_workbook(path)
        try:
            assert wb.active["B2"].value == "Alice"
            assert wb.active["C2"].value == IN_LOCKER_TOKEN
        finally:
            wb.close()

    def test_maintenance_does_not_write_in_locker_token(self, db_session, tmp_path):
        """MAINTENANCE Location is left as-is, not rewritten as Locker."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "Workshop"],
        ])
        DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
            status=DeviceStatus.MAINTENANCE.value,
        )
        db_session.flush()
        write_location(db_session, path)
        assert _location_by_pm(path)["PM-001"] == "Workshop"

    def test_pm_case_and_excel_float_join(self, db_session, tmp_path):
        """Write-back matches PM-001 vs pm-001 and Excel 1001.0 vs SQLite 1001."""
        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["pm-001", "Old"],
            [1001.0, "Old2"],
        ])
        DeviceRepository.create(
            db_session,
            name="A",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        DeviceRepository.create(
            db_session,
            name="B",
            device_type="general",
            pm_number="1001",
            locker_slot=2,
        )
        db_session.flush()
        write_location(db_session, path)
        cells = _location_by_pm(path)
        assert cells["pm-001"] == IN_LOCKER_TOKEN
        assert IN_LOCKER_TOKEN in cells.values()


class TestWritebackOffRequestPath:
    """I10: HTTP/NFC borrow-return return before Excel I/O."""

    def test_borrow_returns_before_excel_io(
        self, db_session, tmp_path, enc_key, hmac_key, monkeypatch
    ):
        """LockerService.borrow returns while write-back is still in a worker."""
        import time

        from smart_locker.auth.session_manager import SessionManager
        from smart_locker.sync import location_writeback as wb
        from smart_locker.sync.location_writeback import (
            WritebackResult,
            flush_scheduled_writeback,
        )

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", IN_LOCKER_TOKEN],
        ])
        monkeypatch.setattr("config.settings.SOURCE_EXCEL_PATH", str(path))
        user = _add_user(db_session, enc_key, hmac_key, "Alice")
        device = DeviceRepository.create(
            db_session,
            name="Meter",
            device_type="general",
            pm_number="PM-001",
            locker_slot=1,
        )
        db_session.flush()
        session = SessionManager(timeout_seconds=60).start_session(user)

        def slow_write(engine, source_path):
            time.sleep(0.4)
            return WritebackResult(written=1, saved=True)

        monkeypatch.setattr(wb, "write_location_with_engine", slow_write)
        t0 = time.monotonic()
        assert LockerService.borrow_device(db_session, session, device.id) is True
        elapsed = time.monotonic() - t0
        assert elapsed < 0.25
        flush_scheduled_writeback()

    def test_writers_serialize(self, tmp_path, monkeypatch):
        """Application writers take one lock; calls do not overlap."""
        import time
        from threading import Thread

        from smart_locker.sync import location_writeback as wb
        from smart_locker.sync.location_writeback import write_location_values

        path = _workbook(tmp_path / "device-list.xlsx", [
            ["Equipment", "Location"],
            ["PM-001", "A"],
            ["PM-002", "B"],
        ])
        order: list[str] = []
        orig = wb._write_once

        def slow(p, wanted):
            order.append("start")
            time.sleep(0.12)
            result = orig(p, wanted)
            order.append("end")
            return result

        monkeypatch.setattr(wb, "_write_once", slow)
        threads = [
            Thread(target=write_location_values, args=(path, {"PM-001": "Alice"})),
            Thread(target=write_location_values, args=(path, {"PM-002": "Bob"})),
        ]
        for t in threads:
            t.start()
        for t in threads:
            t.join()
        assert order == ["start", "end", "start", "end"]

