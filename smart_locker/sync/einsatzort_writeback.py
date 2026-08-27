"""
File: einsatzort_writeback.py
Description: Pi → Excel write-back of Aktueller Einsatzort only, matched by PM.
             Locker available (or any non-borrowed state) writes a stable
             in-locker token; borrowed writes the borrower's display name.
             Other columns and sheets are left untouched. Never inserts rows.
Project: smart_locker/sync
Notes: Called after borrow/return, Register Device, and source import.
       A locked or missing workbook is logged and skipped — the kiosk and
       SQLite stay correct; the next sync retries. Unchanged cells skip
       the save so a local file-watcher cannot loop. If the sheet changes
       between copy and replace, the write is retried from the latest file
       so catalog edits are not reverted.
"""

from __future__ import annotations

import logging
import os
import shutil
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from smart_locker.database.models import Device, DeviceStatus
from smart_locker.database.repositories import DeviceRepository
from smart_locker.sync.source_import import LOCATION_CANDIDATES, PM_CANDIDATES, find_column

logger = logging.getLogger(__name__)

# Stable token written when a locker device is in the cabinet (not borrowed).
# Must contain "schrank" so a later catalog import does not treat it as a person.
IN_LOCKER_TOKEN = "Schrank"

_MAX_RETRIES = 3
_RETRY_DELAY_SECONDS = 1.0


class _StaleWorkbook(Exception):
    """The source xlsx changed on disk after we copied it; retry from the latest file."""


@dataclass
class WritebackResult:
    """Outcome of one Einsatzort write-back. Never used to fail the kiosk."""

    written: int = 0
    unchanged: int = 0
    skipped: int = 0
    saved: bool = False


def _einsatzort_value(device: Device) -> str:
    """Map one locker device to the Excel location cell.

    Args:
        device: Locker row (status + optional borrower already loaded).

    Returns:
        Borrower display name when borrowed, otherwise ``IN_LOCKER_TOKEN``.
    """
    if device.status == DeviceStatus.BORROWED and device.current_borrower is not None:
        name = (device.current_borrower.display_name or "").strip()
        if name:
            return name
    return IN_LOCKER_TOKEN


def _cell_text(value) -> str:
    """Normalize an Excel cell to a stripped string.

    Args:
        value: Raw openpyxl cell value.

    Returns:
        Stripped text, or empty string when the cell is empty.
    """
    if value is None:
        return ""
    return str(value).strip()


def _wanted_by_pm(session: Session) -> dict[str, str]:
    """Build PM → Einsatzort text for every locker device.

    Args:
        session: Active database session (autoflush sees in-request borrows).

    Returns:
        Mapping of stripped PM number to the location string to write.
    """
    wanted: dict[str, str] = {}
    for device in DeviceRepository.list_all(session):
        pm = (device.pm_number or "").strip()
        if not pm:
            continue
        wanted[pm] = _einsatzort_value(device)
    return wanted


def _replace_into(tmp_path: Path, dest: Path) -> None:
    """Atomically replace ``dest`` with ``tmp_path``, retrying a lock.

    Args:
        tmp_path: Staged workbook in the same directory as ``dest``.
        dest: Target ``device-list.xlsx`` on the share or disk.

    Raises:
        PermissionError: Still locked after ``_MAX_RETRIES`` attempts.
        OSError: Replace failed for a reason other than a lock.
    """
    for attempt in range(1, _MAX_RETRIES + 1):
        try:
            tmp_path.replace(dest)
            return
        except PermissionError:
            if attempt < _MAX_RETRIES:
                logger.debug(
                    "Einsatzort write-back: %s is locked, retrying in %ss "
                    "(attempt %d/%d)",
                    dest, _RETRY_DELAY_SECONDS, attempt, _MAX_RETRIES,
                )
                time.sleep(_RETRY_DELAY_SECONDS)
            else:
                raise


def _write_once(session: Session, path: Path) -> WritebackResult:
    """Copy, edit Aktueller Einsatzort, and replace if anything changed.

    Args:
        session: Database session used to read locker devices.
        path: Company ``device-list.xlsx``.

    Returns:
        Counts of written / unchanged / skipped PMs and whether a save ran.

    Raises:
        OSError: Copy, load, save, or replace failed (caller retries/logs).
    """
    result = WritebackResult()
    wanted = _wanted_by_pm(session)
    if not wanted:
        return result

    mtime = path.stat().st_mtime
    work_fd, work_str = tempfile.mkstemp(suffix=".xlsx")
    os.close(work_fd)
    work_path = Path(work_str)
    dest_path: Path | None = None
    wb = None
    try:
        shutil.copy2(path, work_path)
        wb = load_workbook(work_path)
        ws = wb.active
        headers = [
            str(cell.value).strip() if cell.value else ""
            for cell in ws[1]
        ]
        pm_idx = find_column(headers, PM_CANDIDATES)
        loc_idx = find_column(headers, LOCATION_CANDIDATES)
        if pm_idx is None:
            logger.warning(
                "Einsatzort write-back skipped — no PM column in %s (headers: %s).",
                path, headers,
            )
            return result
        if loc_idx is None:
            logger.warning(
                "Einsatzort write-back skipped — no Aktueller Einsatzort column "
                "in %s (headers: %s).",
                path, headers,
            )
            return result

        pm_col = pm_idx + 1
        loc_col = loc_idx + 1
        seen: set[str] = set()
        dirty = False
        for row_i in range(2, ws.max_row + 1):
            pm = _cell_text(ws.cell(row=row_i, column=pm_col).value)
            if not pm or pm not in wanted:
                continue
            seen.add(pm)
            new_value = wanted[pm]
            cell = ws.cell(row=row_i, column=loc_col)
            if _cell_text(cell.value) == new_value:
                result.unchanged += 1
                continue
            cell.value = new_value
            result.written += 1
            dirty = True

        result.skipped = len(wanted) - len(seen)
        if not dirty:
            return result

        dest_fd, dest_str = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
        os.close(dest_fd)
        dest_path = Path(dest_str)
        wb.save(dest_path)
        wb.close()
        wb = None
        try:
            if path.stat().st_mtime != mtime:
                raise _StaleWorkbook()
        except FileNotFoundError:
            raise _StaleWorkbook() from None
        _replace_into(dest_path, path)
        dest_path = None
        result.saved = True
        logger.info(
            "Einsatzort write-back: %d written, %d unchanged, %d not in Excel (%s).",
            result.written, result.unchanged, result.skipped, path,
        )
        return result
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
        work_path.unlink(missing_ok=True)
        if dest_path is not None:
            dest_path.unlink(missing_ok=True)


def write_einsatzort(session: Session, source_path: str | Path) -> WritebackResult:
    """Write locker location into Aktueller Einsatzort. Never raises.

    Args:
        session: Active database session (reads devices; does not commit).
        source_path: Path to the company ``device-list.xlsx``.

    Returns:
        WritebackResult. ``saved`` is False when nothing changed or the
        file could not be written.
    """
    result = WritebackResult()
    path = Path(source_path) if source_path else None
    if path is None or not str(source_path).strip():
        return result
    if not path.exists():
        logger.warning("Einsatzort write-back skipped — file not found: %s", path)
        return result

    try:
        for attempt in range(1, _MAX_RETRIES + 1):
            try:
                return _write_once(session, path)
            except _StaleWorkbook:
                if attempt < _MAX_RETRIES:
                    logger.info(
                        "Einsatzort write-back: %s changed during edit, "
                        "retrying (%d/%d).",
                        path, attempt, _MAX_RETRIES,
                    )
                    continue
                logger.warning(
                    "Einsatzort write-back skipped — %s kept changing during edit.",
                    path,
                )
                return result
    except PermissionError:
        logger.warning(
            "Einsatzort write-back skipped — %s is locked (open in Excel).",
            path,
        )
        return result
    except OSError as e:
        logger.warning(
            "Einsatzort write-back skipped — %s unavailable or unwritable (%s).",
            path, e,
        )
        return result
    except Exception:
        logger.exception(
            "Einsatzort write-back failed for %s — locker database is unchanged.",
            path,
        )
        return result


def write_einsatzort_with_engine(engine, source_path: str | Path) -> WritebackResult:
    """Write Einsatzort using a short-lived session on ``engine``. Never raises.

    Args:
        engine: SQLAlchemy engine (committed locker state).
        source_path: Path to ``device-list.xlsx``.

    Returns:
        WritebackResult from ``write_einsatzort``.
    """
    try:
        with Session(engine) as session:
            return write_einsatzort(session, source_path)
    except Exception:
        logger.exception(
            "Einsatzort write-back failed for %s — locker database is unchanged.",
            source_path,
        )
        return WritebackResult()


def maybe_write_einsatzort(session: Session) -> None:
    """Write-back when ``SOURCE_EXCEL_PATH`` is set. Never raises.

    Args:
        session: The request/tap session that just changed borrow state.

    Returns:
        None.
    """
    try:
        import config.settings as settings

        path = settings.SOURCE_EXCEL_PATH
        if not path:
            return
        write_einsatzort(session, path)
    except Exception:
        logger.exception(
            "Einsatzort write-back failed — locker database is unchanged."
        )
