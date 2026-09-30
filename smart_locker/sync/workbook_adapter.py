"""
File: workbook_adapter.py
Description: Workbook port for copied reads and atomic active-sheet edits.
Project: smart_locker/sync
Notes: Keeps openpyxl and workbook filesystem handling out of sync policy.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
import time
from collections.abc import Callable
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

WRITE_RETRY_ATTEMPTS = 3
_RETRY_DELAY_SECONDS = 1.0


class WorkbookStaleError(Exception):
    """The source workbook changed after its edit copy was made."""


def _file_digest(path: Path) -> str:
    """SHA-256 hex digest of a file's bytes, streamed."""
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _file_identity(path: Path) -> tuple | None:
    """Stat identity (dev, ino, mtime_ns, size), or None when missing."""
    try:
        st = path.stat()
    except FileNotFoundError:
        return None
    return (st.st_dev, st.st_ino, st.st_mtime_ns, st.st_size)


def _catalog_worksheet(wb):
    """Pick the sheet whose header row carries a PM column, else the active one.

    The mirror rewrites the catalog sheet; a workbook left open on a scratch
    sheet (Notes, a pivot, …) must not have that sheet wiped instead.
    """
    from smart_locker.sync.catalog_sheet import find_column, pm_candidates

    for ws in wb.worksheets:
        first_row = next(
            ws.iter_rows(min_row=1, max_row=1, values_only=True), None
        )
        if not first_row:
            continue
        headers = [str(c).strip() if c is not None else "" for c in first_row]
        if find_column(headers, pm_candidates()) is not None:
            return ws
    return wb.active


@dataclass(frozen=True)
class WorkbookRows:
    """Rows read from one workbook sheet, or the reason they were unavailable."""

    rows: list | None
    error: str | None = None


class WorkbookAdapter:
    """Open, read, and edit one xlsx workbook without catalog policy."""

    def __init__(self, path: str | Path) -> None:
        """Store the workbook path.

        Args:
            path: Source ``.xlsx`` file.
        """
        self.path = Path(path)
        self.last_written_digest: str | None = None

    def read_rows(self, sheet_name: str | None = None) -> WorkbookRows:
        """Read values from a copied workbook so an open share file can be read.

        Args:
            sheet_name: Sheet to read, or None for the catalog sheet — the
                same sheet the writer picks, so a workbook left open on a
                scratch tab (Notes, a pivot, …) still reads its catalog.

        Returns:
            WorkbookRows with all cell values or an existing user-facing error.
        """
        if not self.path.exists():
            return WorkbookRows(None, f"File not found: {self.path}")

        try:
            with self._copied_workbook() as work_path:
                try:
                    wb = load_workbook(work_path, read_only=True, data_only=True)
                    try:
                        if sheet_name:
                            if sheet_name not in wb.sheetnames:
                                return WorkbookRows(
                                    None,
                                    f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}",
                                )
                            ws = wb[sheet_name]
                        else:
                            ws = _catalog_worksheet(wb)
                        return WorkbookRows(list(ws.iter_rows(values_only=True)))
                    finally:
                        wb.close()
                except (OSError, BadZipFile, InvalidFileException) as exc:
                    return WorkbookRows(None, f"Source workbook unreadable: {exc}")
        except PermissionError:
            return WorkbookRows(None, f"Source file locked: {self.path}")
        except OSError as exc:
            return WorkbookRows(None, f"Source file unavailable: {self.path} ({exc})")

    def content_digest(self) -> str:
        """SHA-256 hex digest of the workbook's current bytes.

        Returns:
            The digest string.

        Raises:
            OSError: The file could not be read.
        """
        return _file_digest(self.path)

    def edit_active_sheet(
        self,
        edit: Callable[[object], bool],
        expected_mtime: float | None = None,
        pick_sheet: Callable[[object], object] | None = None,
        expected_digest: str | None = None,
    ) -> bool:
        """Copy, edit, and atomically replace the workbook when ``edit`` changes it.

        The replace is bound to the exact bytes that were copied: identity
        (dev/ino/mtime_ns/size) and SHA-256 are re-compared right before every
        replace attempt, so an external edit — even one that restores mtime —
        raises instead of being silently overwritten. That check-to-replace
        gap is inherently small; nothing here pretends to lock against Excel.

        Args:
            edit: Receives the chosen worksheet and returns True when it changed.
            expected_mtime: When given, the current mtime must match — binds
                the copy to an earlier detection stat.
            pick_sheet: Optional chooser ``wb -> worksheet``; default is the
                active sheet.
            expected_digest: When given, the copied bytes must carry this
                SHA-256 — binds the copy to an earlier detection digest.

        Returns:
            True when a changed workbook was saved and replaced; False when
            ``edit`` made no changes.

        Raises:
            WorkbookStaleError: The source changed after it was copied, or
                does not match ``expected_mtime``/``expected_digest``.
            OSError: The workbook could not be copied, loaded, saved, or replaced.
        """
        mtime = self.path.stat().st_mtime
        if expected_mtime is not None and mtime != expected_mtime:
            raise WorkbookStaleError()
        identity = _file_identity(self.path)
        dest_path: Path | None = None
        wb = None
        try:
            with self._copied_workbook() as work_path:
                copy_digest = _file_digest(work_path)
                if expected_digest is not None and copy_digest != expected_digest:
                    raise WorkbookStaleError()
                wb = load_workbook(work_path)
                ws = pick_sheet(wb) if pick_sheet is not None else wb.active
                if not edit(ws):
                    return False

                dest_fd, dest_str = tempfile.mkstemp(
                    suffix=".xlsx", dir=self.path.parent
                )
                os.close(dest_fd)
                dest_path = Path(dest_str)
                wb.save(dest_path)
                wb.close()
                wb = None
                staged_digest = _file_digest(dest_path)
                self._replace_into(
                    dest_path,
                    expected_digest=copy_digest,
                    expected_identity=identity,
                )
                self.last_written_digest = staged_digest
                dest_path = None
                return True
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass
            if dest_path is not None:
                dest_path.unlink(missing_ok=True)

    def write_sheet(
        self,
        headers: list,
        rows: list[list],
        expected_mtime: float | None = None,
        expected_digest: str | None = None,
    ) -> bool:
        """Rewrite the catalog sheet to exactly ``headers`` + ``rows``.

        The mirror owns this file: the sheet carrying the PM column is
        cleared and rewritten (the active sheet is only the fallback);
        other sheets are preserved. A missing file is created — but a
        missing parent directory raises instead of being created, so a
        write can never plant a shadow file under an unmounted share.

        Args:
            headers: Header row cell values.
            rows: Data rows; each a list of cell values (str/date/None).
            expected_mtime: When given, the file's current mtime must match —
                a changed file is refused so a hand edit is never rewritten
                unreviewed.
            expected_digest: When given, the file's current SHA-256 must match
                — a changed file with an unchanged mtime is refused too.

        Returns:
            True when the workbook was saved and replaced.

        Raises:
            WorkbookStaleError: The source changed after it was copied, does
                not match ``expected_mtime``/``expected_digest``, or appeared
                during a create write.
            OSError: The workbook could not be copied, loaded, saved, or replaced.
        """
        from smart_locker.sync.catalog_sheet import safe_cell_value

        if expected_mtime is not None:
            try:
                if self.path.stat().st_mtime != expected_mtime:
                    raise WorkbookStaleError()
            except FileNotFoundError:
                raise WorkbookStaleError() from None

        def _fill(ws) -> None:
            ws.append([safe_cell_value(h) for h in headers])
            for row in rows:
                ws.append([safe_cell_value(v) for v in row])

        if not self.path.exists():
            # No mkdir: a missing parent means the share is not mounted —
            # creating it would plant a shadow file the mount then hides.
            from openpyxl import Workbook

            wb = Workbook()
            _fill(wb.active)
            dest_fd, dest_str = tempfile.mkstemp(
                suffix=".xlsx", dir=self.path.parent
            )
            os.close(dest_fd)
            dest_path = Path(dest_str)
            try:
                wb.save(dest_path)
                wb.close()
                staged_digest = _file_digest(dest_path)
                self._replace_into(dest_path, expected_absent=True)
                self.last_written_digest = staged_digest
                dest_path = None
                return True
            finally:
                if dest_path is not None:
                    dest_path.unlink(missing_ok=True)

        def edit(ws) -> bool:
            if ws.max_row:
                ws.delete_rows(1, ws.max_row)
            _fill(ws)
            return True

        return self.edit_active_sheet(
            edit,
            expected_mtime=expected_mtime,
            pick_sheet=_catalog_worksheet,
            expected_digest=expected_digest,
        )

    def _replace_into(
        self,
        staged_path: Path,
        *,
        expected_digest: str | None = None,
        expected_identity: tuple | None = None,
        expected_absent: bool = False,
    ) -> None:
        """Replace the source workbook with a staged file, retrying a file lock.

        The source is re-verified before every attempt — a write that slept
        on a PermissionError cannot clobber a file that changed (or, for
        ``expected_absent``, appeared) during the delay.

        Args:
            staged_path: The new file to swap in.
            expected_digest: When given, the current file's SHA-256 must match.
            expected_identity: When given, the current file's stat identity
                must match.
            expected_absent: When True, the target must still not exist.

        Raises:
            WorkbookStaleError: A guard did not hold.
            OSError/PermissionError: The replace itself failed.
        """
        for attempt in range(1, WRITE_RETRY_ATTEMPTS + 1):
            if expected_absent and self.path.exists():
                raise WorkbookStaleError()
            if expected_identity is not None and _file_identity(self.path) != expected_identity:
                raise WorkbookStaleError()
            if expected_digest is not None and _file_digest(self.path) != expected_digest:
                raise WorkbookStaleError()
            try:
                staged_path.replace(self.path)
                return
            except PermissionError:
                if attempt == WRITE_RETRY_ATTEMPTS:
                    raise
                time.sleep(_RETRY_DELAY_SECONDS)

    @contextmanager
    def _copied_workbook(self):
        """Yield a temporary copy of the source workbook and delete it afterward."""
        work_fd, work_str = tempfile.mkstemp(suffix=".xlsx")
        os.close(work_fd)
        work_path = Path(work_str)
        try:
            shutil.copy2(self.path, work_path)
            yield work_path
        finally:
            work_path.unlink(missing_ok=True)
